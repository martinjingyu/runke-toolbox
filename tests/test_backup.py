from pathlib import Path

import openpyxl
import pytest

from core.backup import atomic_save_with_backup


def _make_wb(value: str):
    wb = openpyxl.Workbook()
    wb.active["A1"] = value
    return wb


def test_atomic_save_with_backup_swaps_in_new_content_and_keeps_old_as_backup(tmp_path):
    path = tmp_path / "target.xlsx"
    _make_wb("old").save(path)

    backup_path = atomic_save_with_backup(_make_wb("new"), path)

    assert openpyxl.load_workbook(path).active["A1"].value == "new"
    assert backup_path.exists()
    assert openpyxl.load_workbook(backup_path).active["A1"].value == "old"
    # 没有残留写入过程中的临时文件
    assert list(tmp_path.glob("*.写入中-*")) == []


def test_atomic_save_with_backup_leaves_original_untouched_when_save_fails(tmp_path):
    path = tmp_path / "target.xlsx"
    _make_wb("old").save(path)

    class _BrokenWorkbook:
        def save(self, _path):
            raise RuntimeError("模拟存盘失败（比如磁盘满）")

    with pytest.raises(RuntimeError):
        atomic_save_with_backup(_BrokenWorkbook(), path)

    # 原文件完全没被碰过，内容还是旧的
    assert openpyxl.load_workbook(path).active["A1"].value == "old"
    # 没有生成任何备份文件（因为压根没换上新内容，不需要"恢复"）
    assert list(tmp_path.glob("*.备份-*")) == []


def test_atomic_save_with_backup_works_when_target_does_not_exist_yet(tmp_path):
    path = tmp_path / "brand_new.xlsx"

    backup_path = atomic_save_with_backup(_make_wb("first"), path)

    assert openpyxl.load_workbook(path).active["A1"].value == "first"
    assert not backup_path.exists()  # 原本就没有旧文件，自然也没有备份可言
