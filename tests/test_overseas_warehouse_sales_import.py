from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from modules.overseas_warehouse.sales_import.header_map import HeaderParseError, parse_sheet_header
from modules.overseas_warehouse.sales_import.planner import PlanError, apply_plan, build_plan
from modules.overseas_warehouse.sales_import.platform_rules import (
    UnknownShopError,
    classify_erp_platform,
    classify_erp_warehouse,
    guess_platform_from_filename,
)
from modules.overseas_warehouse.sales_import.source_import import (
    SourceReadError,
    SourceRecord,
    aggregate_records,
    read_castlegate_csv,
    read_erp_export,
)
from modules.overseas_warehouse.sales_import.xlsx_writer import _patch_sheet_xml, apply_cell_updates

REAL_DATA_DIR = Path(r"\\Rk\公共文件\个人\黄靖禺\海外仓-刘彩云")


# ---------------------------------------------------------------------------
# 造一张跟真实汇总表同款三层表头的 sheet：
#   col1 = RK-SKU, col2 = 品名
#   block1（col3 汇总列 + col4-6 是 1/2/3 号）：WF-RX，row1 的合并单元格文字可变
#   block2（col7 汇总列 + col8-10）：OS
#   block3（col11 汇总列 + col12-14）：WF-RQ（用"Wayfair（RQ）"这种带中文括号的写法）
#   block4（col15 汇总列 + col16-18）：WF-TS（用"TS"这种缩写）
# ---------------------------------------------------------------------------

_BLOCKS = [
    (None, (4, 6)),  # 第一个区块位置固定是 WF-RX，label 故意留空，模拟真实文件里的空白
    ("OS", (8, 10)),
    ("Wayfair（RQ）", (12, 14)),
    ("TS", (16, 18)),
]


def _build_sheet(wb, name, skus, first_block_label=None):
    ws = wb.create_sheet(name)
    ws.cell(row=3, column=1, value="RK-SKU")
    ws.cell(row=3, column=2, value="品名")

    blocks = list(_BLOCKS)
    if first_block_label is not None:
        blocks[0] = (first_block_label, blocks[0][1])

    for label, (day_start_col, day_end_col) in blocks:
        total_col = day_start_col - 1
        ws.cell(row=3, column=total_col, value="xxx Single SKU\n销售总量")
        # 真实文件里，就算第一个区块没有平台名文字，第1行照样是"合并了但是空白/一堆
        # 空格"的状态，不是压根没合并——这里必须照抄这个状态，不然 header_map 根本
        # 发现不了这个区块的存在（这个坑是先这么写漏了才发现的，不是凭空加的）。
        ws.merge_cells(start_row=1, start_column=day_start_col, end_row=1, end_column=day_end_col)
        if label is not None:
            ws.cell(row=1, column=day_start_col, value=label)
        for i, col in enumerate(range(day_start_col, day_end_col + 1), start=1):
            ws.cell(row=3, column=col, value=f"{i}号")

    for i, (sku, product_name) in enumerate(skus, start=4):
        ws.cell(row=i, column=1, value=sku)
        ws.cell(row=i, column=2, value=product_name)

    return ws


def _build_target_workbook(tmp_path) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _build_sheet(wb, "TX1", [("TD-1", "台灯A"), ("TD-2", "台灯B")])
    _build_sheet(wb, "IL", [("TD-3", "台灯C")])
    _build_sheet(wb, "CA1 ", [("TD-4", "台灯D")])
    _build_sheet(wb, "CG", [("TD-5", "台灯E"), ("TD-6", "台灯F")])
    path = tmp_path / "target.xlsx"
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# header_map
# ---------------------------------------------------------------------------

def test_parse_sheet_header_locates_platform_blocks_and_sku_rows(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _build_sheet(wb, "TX1", [("TD-1", "台灯A"), ("TD-2", "台灯B")])

    header = parse_sheet_header(wb["TX1"])

    assert header.sku_column == 1
    assert header.sku_rows == {"TD-1": 4, "TD-2": 5}
    assert set(header.platforms) == {"WF-RX", "OS", "WF-RQ", "WF-TS"}
    assert header.platforms["WF-RX"].day_columns == {1: 4, 2: 5, 3: 6}
    assert header.platforms["OS"].day_columns == {1: 8, 2: 9, 3: 10}
    assert header.platforms["WF-RQ"].day_columns == {1: 12, 2: 13, 3: 14}
    assert header.platforms["WF-TS"].day_columns == {1: 16, 2: 17, 3: 18}


def test_parse_sheet_header_first_block_is_always_wf_rx_regardless_of_label_text(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _build_sheet(wb, "CG", [("TD-1", "台灯A")], first_block_label="Wayfair All SKU")

    header = parse_sheet_header(wb["CG"])

    assert "WF-RX" in header.platforms
    assert header.platforms["WF-RX"].day_columns == {1: 4, 2: 5, 3: 6}


def test_parse_sheet_header_rejects_unknown_platform_label():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("X")
    ws.cell(row=3, column=1, value="RK-SKU")
    # 第一个区块（不管文字是什么，位置上固定当 WF-RX）+ 第二个区块用一个不认识的平台名，
    # 这样才是真的在测"第二个及以后的区块，平台名对不上就要报错"，不会被"第一个区块
    # 固定是 WF-RX"这条特例挡住。
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=3)
    ws.cell(row=3, column=2, value="1号")
    ws.cell(row=3, column=3, value="2号")
    ws.merge_cells(start_row=1, start_column=5, end_row=1, end_column=6)
    ws.cell(row=1, column=5, value="不认识的平台")
    ws.cell(row=3, column=5, value="1号")
    ws.cell(row=3, column=6, value="2号")

    with pytest.raises(HeaderParseError, match="不在已知的平台对照表"):
        parse_sheet_header(ws)


def test_parse_sheet_header_handles_row1_merge_that_also_covers_the_total_column():
    # 真实文件里的坑：大部分区块第1行的合并单元格只盖住31个日期列，但 CA1 sheet 的 HD
    # 区块，合并单元格连"该平台汇总"那一列（区块最左边那一列）也一起盖了进去（32列宽）。
    # 日期列范围必须按 header_row 上连续的"N号"自己认，不能假设等于第1行合并单元格宽度。
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("X")
    ws.cell(row=3, column=1, value="RK-SKU")
    # 第一个区块：WF-RX，正常宽度（合并只盖3个日期列）
    ws.merge_cells(start_row=1, start_column=3, end_row=1, end_column=5)
    ws.cell(row=3, column=2, value="wayfair Single SKU\n销售总量")
    ws.cell(row=3, column=3, value="1号")
    ws.cell(row=3, column=4, value="2号")
    ws.cell(row=3, column=5, value="3号")
    # 第二个区块：HD，合并单元格从"汇总"那一列（col6）就开始了，比日期列范围（col7-9）宽一列
    ws.merge_cells(start_row=1, start_column=6, end_row=1, end_column=9)
    ws.cell(row=1, column=6, value="HD")
    ws.cell(row=3, column=6, value="HD单SKU\n销售总量")
    ws.cell(row=3, column=7, value="1号")
    ws.cell(row=3, column=8, value="2号")
    ws.cell(row=3, column=9, value="3号")

    header = parse_sheet_header(ws)

    assert header.platforms["HD"].day_columns == {1: 7, 2: 8, 3: 9}


def test_parse_sheet_header_rejects_duplicate_sku_rows():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = _build_sheet(wb, "TX1", [("TD-1", "台灯A"), ("TD-1", "台灯A-重复")])

    with pytest.raises(HeaderParseError, match="不止一行"):
        parse_sheet_header(ws)


# ---------------------------------------------------------------------------
# platform_rules
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    "raw, expected",
    [
        ("至美通CA1", "CA1"),
        ("至美通IL1", "IL"),
        ("至美通TX1", "TX1"),
        ("TS-WH13", "CG"),
        ("TS-WH25", "CG"),
        ("随便什么没见过的仓库", "CG"),
    ],
)
def test_classify_erp_warehouse(raw, expected):
    assert classify_erp_warehouse(raw) == expected


def test_classify_erp_platform_known_and_unknown():
    assert classify_erp_platform("Overstock") == "OS"
    assert classify_erp_platform("Wayfair US") == "WF-RX"
    assert classify_erp_platform("Wayfair3") == "WF-TS"
    assert classify_erp_platform("Wayfair RQ") == "WF-RQ"
    with pytest.raises(UnknownShopError):
        classify_erp_platform("不认识的店铺")


def test_guess_platform_from_filename():
    assert guess_platform_from_filename("CG-RQ CastleGate_SC_Export.csv") == "WF-RQ"
    assert guess_platform_from_filename("CG-RX CastleGate_SC_Export.csv") == "WF-RX"
    assert guess_platform_from_filename("CG-TS CastleGate_SC_Export.csv") == "WF-TS"
    assert guess_platform_from_filename("随便起的名字.csv") is None


# ---------------------------------------------------------------------------
# source_import
# ---------------------------------------------------------------------------

def _build_erp_export(tmp_path, rows, extension=".xls") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["订单号", "店铺", "发货仓库", "采购SKU", "采购SKU个数"])
    for row in rows:
        ws.append(row)
    path = tmp_path / f"erp{extension}"
    wb.save(path)
    return path


def test_read_erp_export_works_even_though_extension_is_xls(tmp_path):
    path = _build_erp_export(
        tmp_path,
        [
            ["CS1", "Overstock", "至美通CA1", "TD-1", 1],
            ["CS2", "Wayfair US", "至美通IL1", "TD-2", 2],
        ],
    )

    result = read_erp_export(path)

    assert result.records == [
        SourceRecord(warehouse="CA1", platform="OS", sku="TD-1", quantity=1, origin="ERP 第2行"),
        SourceRecord(warehouse="IL", platform="WF-RX", sku="TD-2", quantity=2, origin="ERP 第3行"),
    ]
    assert result.skipped == []


def test_read_erp_export_forces_cg_rows_to_os_regardless_of_shop_field(tmp_path):
    # 业务确认过的规则：CG 自己的分平台数据是靠 CastleGate 那 3 份 CSV 导入的，ERP 里
    # 判定成 CG 仓库的行，不管「店铺」字段实际写的是什么（Wayfair RQ/Wayfair3/...），
    # 统一都算 OS——这条规则只对 CG 生效，TX1/IL/CA1 仍然按「店铺」字段分平台。
    path = _build_erp_export(
        tmp_path,
        [
            ["CS1", "Wayfair RQ", "TS-WH13", "TD-1", 1],
            ["CS2", "Wayfair3", "TS-WH25", "TD-2", 2],
            ["CS3", "Wayfair US", "TS-WH24", "TD-3", 3],
            ["CS4", "Wayfair US", "至美通TX1", "TD-4", 4],  # 非 CG 仓库，平台照旧按店铺字段来
        ],
    )

    result = read_erp_export(path)

    assert result.records == [
        SourceRecord(warehouse="CG", platform="OS", sku="TD-1", quantity=1, origin="ERP 第2行"),
        SourceRecord(warehouse="CG", platform="OS", sku="TD-2", quantity=2, origin="ERP 第3行"),
        SourceRecord(warehouse="CG", platform="OS", sku="TD-3", quantity=3, origin="ERP 第4行"),
        SourceRecord(warehouse="TX1", platform="WF-RX", sku="TD-4", quantity=4, origin="ERP 第5行"),
    ]


def test_read_erp_export_skips_row_with_blank_sku_instead_of_failing_whole_file(tmp_path):
    path = _build_erp_export(
        tmp_path,
        [
            ["CS1", "Overstock", "至美通CA1", "", 1],  # 采购SKU 空白——真实 CastleGate 导出里见过这种情况
            ["CS2", "Wayfair RQ", "至美通IL1", "TD-2", 2],
        ],
    )

    result = read_erp_export(path)

    assert result.records == [
        SourceRecord(warehouse="IL", platform="WF-RQ", sku="TD-2", quantity=2, origin="ERP 第3行"),
    ]
    assert len(result.skipped) == 1
    assert "第2行" in result.skipped[0]


def test_read_erp_export_reports_unknown_shop_with_row_number(tmp_path):
    path = _build_erp_export(tmp_path, [["CS1", "不认识的店铺", "至美通CA1", "TD-1", 1]])

    with pytest.raises(SourceReadError, match="ERP 第2行"):
        read_erp_export(path)


def test_read_erp_export_missing_required_column(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["订单号", "店铺"])  # 缺 发货仓库/采购SKU/采购SKU个数
    path = tmp_path / "erp.xls"
    wb.save(path)

    with pytest.raises(SourceReadError, match="缺少必须的表头列"):
        read_erp_export(path)


def _write_csv(tmp_path, name, header, rows) -> Path:
    import csv

    path = tmp_path / name
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(header)
        writer.writerows(rows)
    return path


def test_read_castlegate_csv_sums_multi_row_quantity(tmp_path):
    path = _write_csv(
        tmp_path,
        "CG-RQ export.csv",
        ["Item Number", "SKU", "Quantity"],
        [["TD-1", "internal-1", "1"], ["TD-1", "internal-2", "2"]],
    )

    result = read_castlegate_csv(path, platform="WF-RQ")

    assert [r.sku for r in result.records] == ["TD-1", "TD-1"]
    assert result.skipped == []
    assert aggregate_records(result.records) == {("CG", "WF-RQ", "TD-1"): 3}


def test_read_castlegate_csv_missing_required_column(tmp_path):
    path = _write_csv(tmp_path, "bad.csv", ["Item Number"], [["TD-1"]])

    with pytest.raises(SourceReadError, match="缺少必须的表头列"):
        read_castlegate_csv(path, platform="WF-RQ")


def test_read_castlegate_csv_skips_row_with_blank_item_number_instead_of_failing_whole_file(tmp_path):
    # 真实数据验证过的情况：CastleGate 导出里偶尔会有一整行订单信息都正常（运单号、价格都有），
    # 但 Item Number/SKU 两个字段偏偏是空的——只跳过这一行，不该让整份文件（可能几百行）导入失败。
    path = _write_csv(
        tmp_path,
        "CG-RX export.csv",
        ["Item Number", "SKU", "Quantity"],
        [["", "", "1"], ["TD-2", "internal-2", "2"]],
    )

    result = read_castlegate_csv(path, platform="WF-RX")

    assert [r.sku for r in result.records] == ["TD-2"]
    assert len(result.skipped) == 1
    assert "第2行" in result.skipped[0]


# ---------------------------------------------------------------------------
# xlsx_writer：这里锁定那个"自闭合格子被当成带值格子、非贪婪匹配吞掉后面一串兄弟格子"
# 的回归——拿真实文件测出来过，必须专门测住不能再犯。
# ---------------------------------------------------------------------------

def test_patch_sheet_xml_does_not_swallow_sibling_self_closed_cells():
    row_xml = (
        '<row r="4">'
        '<c r="A4" s="1"><v>hello</v></c>'
        '<c r="B4" s="2"/>'
        '<c r="C4" s="2"/>'
        '<c r="D4" s="2"/>'
        '<c r="E4" s="3"><f>SUM(A4:D4)</f><v>0</v></c>'
        "</row>"
    )
    sheet_xml = f'<?xml version="1.0"?><worksheet><sheetData>{row_xml}</sheetData></worksheet>'.encode("utf-8")

    patched = _patch_sheet_xml(sheet_xml, {4: {3: 999}}).decode("utf-8")  # 列3 = C

    assert '<c r="C4" s="2"><v>999</v></c>' in patched
    # B4/D4 这两个兄弟自闭合格子必须还在，不能被吞掉
    assert '<c r="B4" s="2"/>' in patched
    assert '<c r="D4" s="2"/>' in patched
    # E4 的公式格子必须完整保留
    assert '<c r="E4" s="3"><f>SUM(A4:D4)</f><v>0</v></c>' in patched


def test_patch_sheet_xml_overwrites_cell_that_already_has_a_value():
    row_xml = '<row r="4"><c r="A4" s="1"><v>1</v></c></row>'
    sheet_xml = f'<?xml version="1.0"?><worksheet><sheetData>{row_xml}</sheetData></worksheet>'.encode("utf-8")

    patched = _patch_sheet_xml(sheet_xml, {4: {1: 42}}).decode("utf-8")

    assert '<c r="A4" s="1"><v>42</v></c>' in patched


def test_patch_sheet_xml_inserts_cell_that_did_not_exist_at_all():
    row_xml = '<row r="4"><c r="A4"><v>1</v></c><c r="C4"><v>3</v></c></row>'
    sheet_xml = f'<?xml version="1.0"?><worksheet><sheetData>{row_xml}</sheetData></worksheet>'.encode("utf-8")

    patched = _patch_sheet_xml(sheet_xml, {4: {2: 2}}).decode("utf-8")  # 列2 = B，原来没有

    assert '<c r="A4"><v>1</v></c><c r="B4"><v>2</v></c><c r="C4"><v>3</v></c>' in patched


def test_apply_cell_updates_only_touches_targeted_sheet_and_cells(tmp_path):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1["A1"] = "old"
    ws1["B1"] = "untouched-neighbor"
    # 强制产生自闭合的 <c s="..."/> 格子：设格式但不设值。
    ws1["C1"].fill = PatternFill(fill_type="solid", start_color="FFFF00", end_color="FFFF00")
    ws2 = wb.create_sheet("Sheet2")
    ws2["A1"] = "sheet2-untouched"
    src = tmp_path / "src.xlsx"
    wb.save(src)

    dest = tmp_path / "dest.xlsx"
    apply_cell_updates(str(src), str(dest), {"Sheet1": {(1, 1): 123}})

    out = openpyxl.load_workbook(dest)
    assert out["Sheet1"]["A1"].value == 123
    assert out["Sheet1"]["B1"].value == "untouched-neighbor"
    assert out["Sheet2"]["A1"].value == "sheet2-untouched"

    import zipfile

    zin = zipfile.ZipFile(src)
    zout = zipfile.ZipFile(dest)
    assert set(zin.namelist()) == set(zout.namelist())
    unchanged = [n for n in zin.namelist() if "sheet1" not in n.lower()]
    for name in unchanged:
        assert zin.read(name) == zout.read(name), f"{name} 不应该被改动"


# ---------------------------------------------------------------------------
# planner：端到端——源数据 -> 改动预览 -> 真的写盘
# ---------------------------------------------------------------------------

def test_build_plan_matches_records_to_correct_cells_and_reports_unmatched(tmp_path):
    target = _build_target_workbook(tmp_path)
    records = [
        SourceRecord(warehouse="TX1", platform="WF-RX", sku="TD-1", quantity=5, origin="test"),
        SourceRecord(warehouse="TX1", platform="WF-RX", sku="TD-1", quantity=3, origin="test2"),  # 同一天同SKU多笔要加总
        SourceRecord(warehouse="CG", platform="OS", sku="TD-999-不存在", quantity=7, origin="test3"),
    ]

    plan = build_plan(target, day=2, records=records)

    assert len(plan.diff_rows) == 1
    d = plan.diff_rows[0]
    assert (d.sheet_name, d.sku, d.platform, d.row, d.col) == ("TX1", "TD-1", "WF-RX", 4, 5)
    assert d.old_value == 0
    assert d.new_value == 8  # 5+3 加总

    assert len(plan.unmatched) == 1
    assert plan.unmatched[0].sku == "TD-999-不存在"
    assert "找不到这个 SKU" in plan.unmatched[0].reason


def test_build_plan_falls_back_to_stripping_last_dash_suffix(tmp_path):
    # 业务确认过的规则：源数据里的 SKU 如果原文对不上，去掉最后一个"-"后面的部分再试
    # 一次（比如"TD-5-RX"对应汇总表里的"TD-5"），匹配上了不算"未匹配"，要单独提示是
    # 靠这个兜底匹配上的，写的时候用汇总表里真实的那个 SKU（"TD-5"）。
    target = _build_target_workbook(tmp_path)
    records = [
        SourceRecord(warehouse="CG", platform="WF-RX", sku="TD-5-RX", quantity=4, origin="CG-RX.csv 第2行"),
    ]

    plan = build_plan(target, day=3, records=records)

    assert plan.unmatched == []
    assert len(plan.diff_rows) == 1
    d = plan.diff_rows[0]
    assert (d.sheet_name, d.sku, d.new_value) == ("CG", "TD-5", 4)

    assert len(plan.fuzzy_matches) == 1
    note = plan.fuzzy_matches[0]
    assert (note.source_sku, note.matched_sku, note.quantity) == ("TD-5-RX", "TD-5", 4)
    assert note.origins == ["CG-RX.csv 第2行"]


def test_build_plan_merges_fuzzy_matched_quantity_into_the_same_row_as_a_direct_match(tmp_path):
    # "TD-5"本身的记录 + 靠兜底匹配上的"TD-5-RX"，两笔都得写进同一行，数量要加在一起，
    # 不能后写的把先写的覆盖掉。
    target = _build_target_workbook(tmp_path)
    records = [
        SourceRecord(warehouse="CG", platform="WF-RX", sku="TD-5", quantity=10, origin="a"),
        SourceRecord(warehouse="CG", platform="WF-RX", sku="TD-5-RX", quantity=4, origin="b"),
    ]

    plan = build_plan(target, day=3, records=records)

    assert len(plan.diff_rows) == 1
    assert plan.diff_rows[0].new_value == 14
    assert len(plan.fuzzy_matches) == 1


def test_build_plan_does_not_over_strip_a_sku_that_only_has_one_dash(tmp_path):
    # "TD-999"这种只有一个"-"、本来就是"货号-序号"基本形态的 SKU，去掉最后一个"-"会变成
    # "TD"，明显是把货号本身切坏了——切完剩下的部分必须还带"-"（还长得像"货号-序号"）
    # 才做这个兜底。这里故意在表里放一个真的叫"TD"的 SKU，如果没有这条守卫，"TD-999"
    # 会被误判成匹配到这个"TD"上去，实际不应该发生。
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _build_sheet(wb, "TX1", [("TD", "不该被匹配上的货号"), ("TD-1", "台灯A")])
    _build_sheet(wb, "IL", [("TD-3", "台灯C")])
    _build_sheet(wb, "CA1 ", [("TD-4", "台灯D")])
    _build_sheet(wb, "CG", [("TD-5", "台灯E")])
    target = tmp_path / "single_dash.xlsx"
    wb.save(target)

    records = [
        SourceRecord(warehouse="TX1", platform="WF-RX", sku="TD-999", quantity=1, origin="test"),
    ]

    plan = build_plan(target, day=1, records=records)

    assert plan.diff_rows == []
    assert plan.fuzzy_matches == []
    assert len(plan.unmatched) == 1
    assert plan.unmatched[0].sku == "TD-999"


def test_build_plan_tolerates_a_missing_warehouse_sheet(tmp_path):
    # 业务确认过：TX1 只放库存，卖完之后这个 sheet 以后会被直接删掉，不是异常状态——
    # 缺了某个仓库的 sheet 不该让整批导入失败，这个仓库的源数据走"未匹配"清单，
    # 其它仓库正常导入。
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _build_sheet(wb, "IL", [("TD-3", "台灯C")])
    _build_sheet(wb, "CA1 ", [("TD-4", "台灯D")])
    _build_sheet(wb, "CG", [("TD-5", "台灯E")])
    target = tmp_path / "no_tx1.xlsx"
    wb.save(target)  # 故意不建 TX1 这个 sheet

    records = [
        SourceRecord(warehouse="TX1", platform="WF-RX", sku="TD-1", quantity=5, origin="test"),
        SourceRecord(warehouse="IL", platform="WF-RX", sku="TD-3", quantity=2, origin="test2"),
    ]

    plan = build_plan(target, day=1, records=records)

    assert len(plan.diff_rows) == 1
    assert plan.diff_rows[0].sheet_name == "IL"

    assert len(plan.unmatched) == 1
    assert plan.unmatched[0].warehouse == "TX1"
    assert "没有「TX1」这个仓库的 sheet" in plan.unmatched[0].reason


def test_build_plan_errors_when_no_warehouse_sheet_exists_at_all(tmp_path):
    # 4 个仓库 sheet 一个都不在，大概率是选错文件了，这种要报错，不能悄悄全塞进未匹配清单。
    wb = openpyxl.Workbook()
    wb.active.title = "跟仓库表毫无关系的sheet"
    target = tmp_path / "wrong_file.xlsx"
    wb.save(target)

    records = [SourceRecord(warehouse="TX1", platform="WF-RX", sku="TD-1", quantity=5, origin="test")]

    with pytest.raises(PlanError, match="是不是选错文件了"):
        build_plan(target, day=1, records=records)


def test_build_plan_rejects_out_of_range_day(tmp_path):
    target = _build_target_workbook(tmp_path)
    with pytest.raises(PlanError, match="1-31"):
        build_plan(target, day=32, records=[])


def test_build_plan_matches_sheets_by_name_not_by_position(tmp_path):
    # 业务提出的顾虑：以后可能会删掉/挪动工作簿里的某些 sheet（比如"销量汇总"那个衍生报表），
    # 4 个仓库 sheet 之间的先后顺序也不一定固定——不能靠 wb.worksheets[i] 这种按位置取 sheet
    # 的写法，得靠 sheet 名字。这里故意把 4 个仓库 sheet 建成跟 _build_target_workbook 完全
    # 相反的顺序，还在最前面插一个不相关、结构也不对的"销量汇总"sheet（如果代码哪里偷偷按
    # 位置取 sheet，这个假的"销量汇总"会第一个被扫到，直接触发表头解析报错），验证还是能
    # 正确工作。
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet("销量汇总").append(["随便什么", "跟仓库表结构完全不一样"])
    _build_sheet(wb, "CG", [("TD-6", "台灯F")])
    _build_sheet(wb, "CA1 ", [("TD-4", "台灯D")])
    _build_sheet(wb, "IL", [("TD-3", "台灯C")])
    _build_sheet(wb, "TX1", [("TD-1", "台灯A")])
    target = tmp_path / "shuffled.xlsx"
    wb.save(target)

    records = [SourceRecord(warehouse="TX1", platform="WF-RX", sku="TD-1", quantity=7, origin="test")]
    plan = build_plan(target, day=1, records=records)

    assert len(plan.diff_rows) == 1
    assert plan.diff_rows[0].sheet_name == "TX1"
    assert plan.diff_rows[0].new_value == 7


def test_apply_plan_writes_values_and_keeps_backup(tmp_path):
    target = _build_target_workbook(tmp_path)
    records = [
        SourceRecord(warehouse="CA1", platform="WF-TS", sku="TD-4", quantity=9, origin="test"),
    ]
    plan = build_plan(target, day=1, records=records)

    backup_path = apply_plan(plan)

    assert backup_path.exists()
    wb_after = openpyxl.load_workbook(target)
    assert wb_after["CA1 "].cell(row=4, column=16).value == 9
    # 备份文件里应该还是改之前的旧值（0/None）
    wb_backup = openpyxl.load_workbook(backup_path)
    assert not wb_backup["CA1 "].cell(row=4, column=16).value


# ---------------------------------------------------------------------------
# 真实数据集成测试：本机（或者能访问公司内网共享）才跑，跑一遍完整流程——真实 ERP 导出 +
# 3 份真实 CastleGate CSV 导出 -> 目标表（用测试专用副本，13号数据已经被业务特意清空方便
# 测试）。锁定的是开发过程中真的从这份数据踩出来的几个坑：ERP 导出后缀是 .xls 但内容其实
# 是 xlsx、CastleGate CSV 里偶尔有整行数据都正常但 Item Number 为空的订单、CA1 sheet 的
# HD 区块第1行合并单元格宽度跟其它区块不一样——都是拿真实文件跑出来才发现的，合成数据
# 造不出来这几个具体的坑，所以额外保留这个真实数据测试。
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not REAL_DATA_DIR.exists(), reason="需要能访问公司内网的海外仓测试数据")
def test_real_sales_import_end_to_end(tmp_path):
    import shutil
    import zipfile

    target = tmp_path / "target.xlsx"
    shutil.copy(REAL_DATA_DIR / "2026年9月US库存销售明细表-Test.xlsx", target)
    erp_path = tmp_path / "erp.xls"
    shutil.copy(REAL_DATA_DIR / "20260913_1706577974866_1789372560376.xls", erp_path)

    records = []
    erp_result = read_erp_export(erp_path)
    records.extend(erp_result.records)
    assert len(erp_result.records) > 0

    csv_files = [
        ("CG-RQ CastleGate_SC_Export_09-14-2026_04-32-15.csv", "WF-RQ"),
        ("CG-RX CastleGate_SC_Export_09-14-2026_04-31-26.csv", "WF-RX"),
        ("CG-TS CastleGate_SC_Export_09-14-2026_04-31-56.csv", "WF-TS"),
    ]
    for name, platform in csv_files:
        result = read_castlegate_csv(REAL_DATA_DIR / name, platform)
        records.extend(result.records)

    plan = build_plan(target, day=13, records=records)
    assert len(plan.diff_rows) > 0

    # 必须在 apply_plan()（内部会把 target 原子改名替换掉）之前就把 ZipFile 关掉——Windows
    # 上一个还开着读句柄的文件没法被 rename 替换，不关掉这个断言本身会把 apply_plan 搞挂。
    with zipfile.ZipFile(target) as zin_before:
        media_before = {n for n in zin_before.namelist() if n.startswith("xl/media/")}
        assert media_before  # 目标表本身真的带着商品图片，不然这个测试没测到重点
        media_bytes_before = {n: zin_before.read(n) for n in media_before}

    apply_plan(plan)

    with zipfile.ZipFile(target) as zin_after:
        media_after = {n for n in zin_after.namelist() if n.startswith("xl/media/")}
        assert media_after == media_before  # 图片一张没少、一张没变
        for name in media_before:
            assert zin_after.read(name) == media_bytes_before[name]  # 字节也完全没变

    wb_after = openpyxl.load_workbook(target, read_only=True, data_only=True)
    for d in plan.diff_rows:
        assert wb_after[d.sheet_name].cell(row=d.row, column=d.col).value == d.new_value
