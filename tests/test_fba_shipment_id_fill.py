import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from modules.logistics.fba_shipment_id_fill.fba_source import FbaSourceError, parse_fba_folder
from modules.logistics.fba_shipment_id_fill.plan_matcher import apply_plan, build_plan

SHIP_DATE = dt.date(2026, 9, 16)


def _write_fba_csv(path: Path, rows: list[list[str]]) -> None:
    header = "货件名称,FBA货件编号,内部编号,目的地,MSKU,预计商品数量,MSKU状态\n"
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        fh.write(header)
        for row in rows:
            fh.write(",".join(str(v) for v in row) + "\n")


# ---------------------------------------------------------------------------
# fba_source
# ---------------------------------------------------------------------------


def test_parse_fba_folder_skips_wrong_date_and_cancelled(tmp_path):
    _write_fba_csv(
        tmp_path / "a.csv",
        [
            ["SX-9.16-SBD1", "FBA1", "IN1", "SBD1", "TD-1", 9, "正在接收"],
            ["SX-9.17-SBD1", "FBA2", "IN2", "SBD1", "TD-2", 9, "正在接收"],  # 日期不对
            ["取消", "FBA3", "IN3", "CLT2", "TD-3", 0, "已取消"],  # 已取消
            ["取消SX-9.16-LAX9", "FBA4", "IN4", "LAX9", "TD-4", 0, "已取消"],  # 已取消+日期能解析
        ],
    )
    result = parse_fba_folder(tmp_path, SHIP_DATE)

    assert list(result.entries_by_msku.keys()) == ["TD-1"]
    assert len(result.skipped) == 3
    reasons = {s.msku: s.reason for s in result.skipped}
    assert "跟筛选日期" in reasons["TD-2"]
    assert "已取消" in reasons["TD-3"]
    assert "已取消" in reasons["TD-4"]


def test_parse_fba_folder_handles_mixed_gbk_and_utf8_files(tmp_path):
    # 实测同一个文件夹里，不同账号爬出来的 CSV 编码不统一：有的带 BOM 是 UTF-8，有的是没有
    # BOM 的 GBK（Excel 另存 CSV 的默认编码）。两种都要能读，不能整份按同一种编码硬读。
    header = "货件名称,FBA货件编号,内部编号,目的地,MSKU,预计商品数量,MSKU状态\n"
    row = "SX-9.16-SBD1,FBA1,IN1,SBD1,TD-GBK,9,正在接收\n"
    with open(tmp_path / "gbk.csv", "w", encoding="gbk", newline="") as fh:
        fh.write(header + row)
    _write_fba_csv(
        tmp_path / "utf8.csv",
        [["SX-9.16-SBD1", "FBA2", "IN2", "SBD1", "TD-UTF8", 9, "正在接收"]],
    )

    result = parse_fba_folder(tmp_path, SHIP_DATE)

    assert set(result.entries_by_msku.keys()) == {"TD-GBK", "TD-UTF8"}


def test_parse_fba_folder_dedups_identical_rows_across_files(tmp_path):
    row = ["SX-9.16-SBD1", "FBA1", "IN1", "SBD1", "TD-1", 9, "正在接收"]
    _write_fba_csv(tmp_path / "a.csv", [row])
    _write_fba_csv(tmp_path / "b.csv", [row])

    result = parse_fba_folder(tmp_path, SHIP_DATE)

    assert len(result.entries_by_msku["TD-1"]) == 1
    assert result.duplicates_removed == 1


def test_parse_fba_folder_zero_quantity_not_cancelled_is_skipped(tmp_path):
    _write_fba_csv(
        tmp_path / "a.csv",
        [["SX-9.16-SBD1", "FBA1", "IN1", "SBD1", "TD-1", 0, "正在接收"]],
    )
    result = parse_fba_folder(tmp_path, SHIP_DATE)

    assert result.entries_by_msku == {}
    assert len(result.skipped) == 1
    assert "数量为 0" in result.skipped[0].reason


def test_parse_fba_folder_missing_header_raises(tmp_path):
    path = tmp_path / "a.csv"
    with open(path, "w", encoding="utf-8-sig") as fh:
        fh.write("货件名称,FBA货件编号\nfoo,bar\n")
    with pytest.raises(FbaSourceError):
        parse_fba_folder(tmp_path, SHIP_DATE)


def test_parse_fba_folder_no_csv_raises(tmp_path):
    with pytest.raises(FbaSourceError):
        parse_fba_folder(tmp_path, SHIP_DATE)


# ---------------------------------------------------------------------------
# plan_matcher：发货计划表构造helper
# ---------------------------------------------------------------------------

PLAN_HEADERS = [
    "采购单号", "型号", "标签", "产品名称", "箱数", "箱容", "数量", "长", "宽", "高",
    "毛重", "CBM", "总材重", "总实重", "仓库", "FBA ID", "追踪编号", "ZD", "编号",
    "备注", "交货时间", "发货时间", "工厂", "货代", "出货单号", "状态",
]


def _make_plan_workbook(rows: list[dict]) -> tuple:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "发货计划(2)"
    for _ in range(4):
        ws.append([])
    ws.append(PLAN_HEADERS)
    col = {name: i + 1 for i, name in enumerate(PLAN_HEADERS)}
    for row in rows:
        r = ws.max_row + 1
        ws.cell(row=r, column=col["采购单号"]).value = row.get("采购单号", "PO")
        ws.cell(row=r, column=col["型号"]).value = row.get("型号", row.get("标签", ""))
        label = row.get("标签", "")
        if row.get("标签_is_formula"):
            ws.cell(row=r, column=col["标签"]).value = f"=+B{r}"
        else:
            ws.cell(row=r, column=col["标签"]).value = label
        ws.cell(row=r, column=col["箱数"]).value = row["箱数"]
        ws.cell(row=r, column=col["箱容"]).value = row.get("箱容", 3)
        ws.cell(row=r, column=col["数量"]).value = f"=E{r}*F{r}"
        ws.cell(row=r, column=col["长"]).value = 100
        ws.cell(row=r, column=col["宽"]).value = 100
        ws.cell(row=r, column=col["高"]).value = 100
        ws.cell(row=r, column=col["FBA ID"]).value = row.get("FBA ID")
        ws.cell(row=r, column=col["ZD"]).value = row.get("ZD", "US")
        ws.cell(row=r, column=col["发货时间"]).value = row.get("发货时间", dt.datetime.combine(SHIP_DATE, dt.time()))
        ws.cell(row=r, column=col["状态"]).value = row.get("状态", "未发货")
    return wb, ws, col


def _fba_entry(msku, boxes_qty, box_capacity, fba_id="FBA1", destination="SBD1", internal_id="IN1"):
    from modules.logistics.fba_shipment_id_fill.fba_source import FbaEntry

    return FbaEntry(
        msku=msku,
        destination=destination,
        fba_id=fba_id,
        internal_id=internal_id,
        quantity_pieces=boxes_qty * box_capacity,
        source_file="a.csv",
        source_row=2,
    )


def _fba_result(entries_by_msku, skipped=None, duplicates_removed=0):
    from modules.logistics.fba_shipment_id_fill.fba_source import FbaSourceResult

    return FbaSourceResult(entries_by_msku=entries_by_msku, skipped=skipped or [], duplicates_removed=duplicates_removed)


# ---------------------------------------------------------------------------
# plan_matcher.build_plan：校验
# ---------------------------------------------------------------------------


def test_build_plan_exact_match_no_split_needed():
    wb, ws, col = _make_plan_workbook([{"标签": "TD-1", "箱数": 5}])
    fba_result = _fba_result({"TD-1": [_fba_entry("TD-1", boxes_qty=5, box_capacity=3, fba_id="FBA1")]})

    plan = build_plan(ws, fba_result, SHIP_DATE)

    assert plan.errors == []
    assert len(plan.row_plans) == 1
    assert len(plan.row_plans[0].pieces) == 1
    assert plan.row_plans[0].pieces[0].fba_id == "FBA1"


def test_build_plan_quantity_mismatch_reports_error():
    wb, ws, col = _make_plan_workbook([{"标签": "TD-1", "箱数": 5}])
    fba_result = _fba_result({"TD-1": [_fba_entry("TD-1", boxes_qty=8, box_capacity=3, fba_id="FBA1")]})

    plan = build_plan(ws, fba_result, SHIP_DATE)

    assert plan.has_blocking_errors
    assert any("箱数对不上" in e for e in plan.errors)
    assert plan.row_plans == []


def test_build_plan_existing_fba_id_reports_error():
    wb, ws, col = _make_plan_workbook([{"标签": "TD-1", "箱数": 5, "FBA ID": "已经有的号"}])
    fba_result = _fba_result({"TD-1": [_fba_entry("TD-1", boxes_qty=5, box_capacity=3)]})

    plan = build_plan(ws, fba_result, SHIP_DATE)

    assert plan.has_blocking_errors
    assert any("已经有 FBA ID" in e for e in plan.errors)


def test_build_plan_missing_matching_row_reports_error():
    wb, ws, col = _make_plan_workbook([{"标签": "TD-OTHER", "箱数": 5}])
    fba_result = _fba_result({"TD-1": [_fba_entry("TD-1", boxes_qty=5, box_capacity=3)]})

    plan = build_plan(ws, fba_result, SHIP_DATE)

    assert plan.has_blocking_errors
    assert any("没有找到" in e for e in plan.errors)


def test_build_plan_not_divisible_reports_error():
    wb, ws, col = _make_plan_workbook([{"标签": "TD-1", "箱数": 5, "箱容": 3}])
    from modules.logistics.fba_shipment_id_fill.fba_source import FbaEntry

    entry = FbaEntry(
        msku="TD-1", destination="SBD1", fba_id="FBA1", internal_id="IN1",
        quantity_pieces=10, source_file="a.csv", source_row=2,
    )
    fba_result = _fba_result({"TD-1": [entry]})

    plan = build_plan(ws, fba_result, SHIP_DATE)

    assert plan.has_blocking_errors
    assert any("除不尽箱容" in e for e in plan.errors)


def test_build_plan_bad_zd_reports_error():
    wb, ws, col = _make_plan_workbook([{"标签": "TD-1", "箱数": 5, "ZD": "MX"}])
    fba_result = _fba_result({"TD-1": [_fba_entry("TD-1", boxes_qty=5, box_capacity=3)]})

    plan = build_plan(ws, fba_result, SHIP_DATE)

    assert plan.has_blocking_errors
    assert any("ZD 不是预期的" in e for e in plan.errors)


def test_build_plan_multiple_problems_all_reported_together():
    wb, ws, col = _make_plan_workbook(
        [
            {"标签": "TD-1", "箱数": 5, "FBA ID": "已存在"},
            {"标签": "TD-2", "箱数": 5, "ZD": "MX"},
        ]
    )
    fba_result = _fba_result(
        {
            "TD-1": [_fba_entry("TD-1", boxes_qty=5, box_capacity=3)],
            "TD-2": [_fba_entry("TD-2", boxes_qty=5, box_capacity=3)],
            "TD-3": [_fba_entry("TD-3", boxes_qty=5, box_capacity=3)],  # 计划表里完全没有
        }
    )

    plan = build_plan(ws, fba_result, SHIP_DATE)

    assert len(plan.errors) == 3


def test_build_plan_ignores_rows_with_different_ship_date():
    wb, ws, col = _make_plan_workbook(
        [{"标签": "TD-1", "箱数": 5, "发货时间": dt.datetime(2026, 1, 1)}]
    )
    fba_result = _fba_result({"TD-1": [_fba_entry("TD-1", boxes_qty=5, box_capacity=3)]})

    plan = build_plan(ws, fba_result, SHIP_DATE)

    # 计划表里没有"发货日期=筛选日期"的匹配行，等价于没有找到
    assert plan.has_blocking_errors
    assert any("没有找到" in e for e in plan.errors)


def test_build_plan_resolves_formula_label():
    wb, ws, col = _make_plan_workbook([{"型号": "TD-1", "标签_is_formula": True, "箱数": 5}])
    fba_result = _fba_result({"TD-1": [_fba_entry("TD-1", boxes_qty=5, box_capacity=3)]})

    plan = build_plan(ws, fba_result, SHIP_DATE)

    assert plan.errors == []
    assert len(plan.row_plans) == 1


# ---------------------------------------------------------------------------
# 拆分算法 + apply_plan
# ---------------------------------------------------------------------------


def test_build_plan_splits_5_5_against_8_2_example_from_user():
    wb, ws, col = _make_plan_workbook(
        [
            {"标签": "TD-1", "箱数": 5, "采购单号": "PO-1"},
            {"标签": "TD-1", "箱数": 5, "采购单号": "PO-2"},
        ]
    )
    fba_result = _fba_result(
        {
            "TD-1": [
                _fba_entry("TD-1", boxes_qty=8, box_capacity=3, fba_id="FBA-A", destination="SBD1"),
                _fba_entry("TD-1", boxes_qty=2, box_capacity=3, fba_id="FBA-B", destination="CLT2"),
            ]
        }
    )

    plan = build_plan(ws, fba_result, SHIP_DATE)
    assert plan.errors == []
    assert len(plan.row_plans) == 2

    row1, row2 = plan.row_plans
    assert [p.boxes for p in row1.pieces] == [5]
    assert row1.pieces[0].fba_id == "FBA-A"
    assert [p.boxes for p in row2.pieces] == [3, 2]
    assert row2.pieces[0].fba_id == "FBA-A"
    assert row2.pieces[1].fba_id == "FBA-B"


def test_apply_plan_writes_fields_and_keeps_split_rows_adjacent():
    wb, ws, col = _make_plan_workbook(
        [
            {"标签": "TD-1", "箱数": 5, "采购单号": "PO-1", "ZD": "US"},
            {"标签": "TD-1", "箱数": 5, "采购单号": "PO-2", "ZD": "US"},
            {"标签": "TD-UNRELATED", "箱数": 7, "采购单号": "PO-3"},
        ]
    )
    unrelated_row = 8  # header 在第5行，数据从第6行开始：6,7,8
    fba_result = _fba_result(
        {
            "TD-1": [
                _fba_entry("TD-1", boxes_qty=8, box_capacity=3, fba_id="FBA-A", destination="SBD1", internal_id="IN-A"),
                _fba_entry("TD-1", boxes_qty=2, box_capacity=3, fba_id="FBA-B", destination="CLT2", internal_id="IN-B"),
            ]
        }
    )

    plan = build_plan(ws, fba_result, SHIP_DATE)
    assert plan.errors == []

    apply_plan(ws, plan)

    # 行6：PO-1，整行不用拆，箱数还是5，FBA-A
    assert ws.cell(row=6, column=col["箱数"]).value == 5
    assert ws.cell(row=6, column=col["FBA ID"]).value == "FBA-A"
    assert ws.cell(row=6, column=col["仓库"]).value == "US(SBD1)"
    assert ws.cell(row=6, column=col["追踪编号"]).value == "IN-A"
    assert ws.cell(row=6, column=col["采购单号"]).value == "PO-1"

    # 行7、8：PO-2 被拆成 3 + 2，紧跟在原来的位置，采购单号保持一致（复制过去的）
    assert ws.cell(row=7, column=col["箱数"]).value == 3
    assert ws.cell(row=7, column=col["FBA ID"]).value == "FBA-A"
    assert ws.cell(row=7, column=col["采购单号"]).value == "PO-2"
    assert ws.cell(row=8, column=col["箱数"]).value == 2
    assert ws.cell(row=8, column=col["FBA ID"]).value == "FBA-B"
    assert ws.cell(row=8, column=col["仓库"]).value == "US(CLT2)"
    assert ws.cell(row=8, column=col["采购单号"]).value == "PO-2"

    # 原来跟这次拆分无关的行被整体挪到新位置（往下移了一行），字段完全没变
    assert ws.cell(row=9, column=col["标签"]).value == "TD-UNRELATED"
    assert ws.cell(row=9, column=col["箱数"]).value == 7
    assert ws.cell(row=9, column=col["采购单号"]).value == "PO-3"
    assert ws.cell(row=9, column=col["FBA ID"]).value is None

    # 数量列公式的自引用行号要跟着新行号走，不能还是指向旧行号
    assert ws.cell(row=8, column=col["数量"]).value == "=E8*F8"
    assert ws.cell(row=9, column=col["数量"]).value == "=E9*F9"


def test_apply_plan_formula_recomputes_correctly_after_split(tmp_path):
    wb, ws, col = _make_plan_workbook([{"标签": "TD-1", "箱数": 10, "箱容": 3}])
    fba_result = _fba_result(
        {
            "TD-1": [
                _fba_entry("TD-1", boxes_qty=6, box_capacity=3, fba_id="FBA-A", destination="SBD1"),
                _fba_entry("TD-1", boxes_qty=4, box_capacity=3, fba_id="FBA-B", destination="CLT2"),
            ]
        }
    )
    plan = build_plan(ws, fba_result, SHIP_DATE)
    assert plan.errors == []
    apply_plan(ws, plan)

    path = tmp_path / "out.xlsx"
    wb.save(path)
    wb2 = openpyxl.load_workbook(path, data_only=False)
    ws2 = wb2["发货计划(2)"]
    from modules.logistics.shipment_plan_apply.column_utils import resolve_cell_value

    assert resolve_cell_value(ws2, 6, col["数量"]) == 6 * 3
    assert resolve_cell_value(ws2, 7, col["数量"]) == 4 * 3


def test_zip_split_one_fba_number_spans_multiple_plan_rows():
    wb, ws, col = _make_plan_workbook(
        [
            {"标签": "TD-1", "箱数": 2, "采购单号": "PO-1"},
            {"标签": "TD-1", "箱数": 2, "采购单号": "PO-2"},
            {"标签": "TD-1", "箱数": 2, "采购单号": "PO-3"},
        ]
    )
    fba_result = _fba_result({"TD-1": [_fba_entry("TD-1", boxes_qty=6, box_capacity=3, fba_id="FBA-ONLY")]})

    plan = build_plan(ws, fba_result, SHIP_DATE)
    assert plan.errors == []
    for row_plan in plan.row_plans:
        assert len(row_plan.pieces) == 1
        assert row_plan.pieces[0].fba_id == "FBA-ONLY"
        assert row_plan.pieces[0].boxes == 2
