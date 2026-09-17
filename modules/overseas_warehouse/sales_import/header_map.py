"""解析海外仓销量汇总表里一个仓库 sheet 的表头，定位"某个平台某一号"对应哪一列、
"某个 SKU"对应哪一行。

表头是 3 层的（第1-3行，数据从第4行开始），而且每个 sheet 的列顺序、列的有无都不一样
（比如 TX1 有单独的「仓库」文本列，CA1 没有，只有「仓库代码」数字列），所以完全不能
按列号硬编码，必须每次都按表头文字重新定位。具体结构见
modules/overseas_warehouse/sales_import 目录下这几个模块的开发过程中拿真实文件验证的
结论（第1行是合并单元格，标真正的平台名；第2行是"该平台全部SKU汇总"的公式展示行，
不是要写的数据；第3行是"该平台单SKU汇总"+"1号"..."31号"的每日销量列，这才是要写的）。

第一个平台区块（离"所有平台销售总量"最近的那个）在几个 sheet 里第1行的合并单元格
文字不统一（有的是空白、有的写"Wayfair"、有的写"Wayfair All SKU"）——这个不是随便
写的，是实际抄真实文件抄出来的规律：不管文字是什么，"第一个区块固定是 WF-RX（Wayfair
US）"这件事跨 4 个 sheet 都成立，所以按位置认第一个区块，其余区块按文字对照表识别，
文字对不上就直接报错，不猜。
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

from openpyxl.worksheet.worksheet import Worksheet

# 平台名对照表：key 是"去空白、去不间断空格、转小写"之后的合并单元格文字。
_PLATFORM_ALIASES = {
    "wayfair": "WF-RX",
    "wayfair all sku": "WF-RX",
    "lowe's": "LOWES",
    "lowes": "LOWES",
    "hd": "HD",
    "walmart": "WALMART",
    "os": "OS",
    "ts": "WF-TS",
    "wayfair（rq）": "WF-RQ",
    "wayfair(rq)": "WF-RQ",
}

_DAY_RE = re.compile(r"(\d{1,2})号\s*$")
_SKU_HEADER_TEXT = "RK-SKU"


class HeaderParseError(Exception):
    """表头长得跟预期不一样，不敢猜，直接报错让人工看。"""


def _norm_label(text) -> str:
    if text is None:
        return ""
    return str(text).replace("\xa0", " ").strip().lower()


@dataclass(frozen=True)
class PlatformBlock:
    platform_key: str
    day_columns: dict[int, int] = field(default_factory=dict)  # 几号 -> 列号(1-based)


@dataclass(frozen=True)
class SheetHeaderMap:
    sheet_name: str
    sku_column: int
    sku_rows: dict[str, int]  # RK-SKU -> 行号(1-based)
    platforms: dict[str, PlatformBlock]  # platform_key -> block


def _find_header_row(ws: Worksheet, max_scan_rows: int = 6) -> tuple[int, int]:
    """扫前几行找到写着 RK-SKU 的那个格子，返回 (表头行号, RK-SKU 列号)。"""
    for row in range(1, max_scan_rows + 1):
        for col in range(1, min(ws.max_column, 40) + 1):
            if ws.cell(row=row, column=col).value == _SKU_HEADER_TEXT:
                return row, col
    raise HeaderParseError(f"sheet {ws.title!r} 前 {max_scan_rows} 行里找不到 {_SKU_HEADER_TEXT!r} 这个表头")


def _parse_day_number(text) -> int | None:
    if text is None:
        return None
    m = _DAY_RE.search(str(text))
    if not m:
        return None
    day = int(m.group(1))
    if 1 <= day <= 31:
        return day
    return None


def _label_at(ws: Worksheet, row: int, col: int) -> object:
    """第1行那个平台名格子，有的 sheet 是合并单元格，有的（万一）不是——不管哪种，
    都按"这个坐标落在哪个合并区域"去找真正的文字，找不到就退回直接读这个格子自己的值。
    """
    for merged_range in ws.merged_cells.ranges:
        if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= col <= merged_range.max_col:
            return ws.cell(row=merged_range.min_row, column=merged_range.min_col).value
    return ws.cell(row=row, column=col).value


def parse_sheet_header(ws: Worksheet) -> SheetHeaderMap:
    header_row, sku_col = _find_header_row(ws)
    platform_label_row = header_row - 2
    if platform_label_row < 1:
        raise HeaderParseError(
            f"sheet {ws.title!r}：表头行是第 {header_row} 行，往上数2行应该是平台名那一行，"
            "但已经数到表格外面了，表头结构跟预期的三层表头不一样"
        )

    # 平台区块的"日期列范围"按 header_row 上连续一串"N号"文字自己认出来，不依赖第1行
    # 合并单元格的宽度——拿真实文件验证过，同一份表里不同平台区块，第1行合并单元格有的
    # 只盖住日期列（31列），有的连"该平台汇总"那一列也一起盖进去了（32列），宽度并不统一，
    # 按合并宽度去切日期列范围会把"汇总"那一列也误当成日期列，进而报错说解析不出是几号。
    # 平台名文字单独按"日期列范围第一列落在第1行哪个合并区域里"去找（见 _label_at），
    # 两件事分开判断，不互相依赖对方的边界。
    day_run_cols: list[int] = []
    for col in range(sku_col + 1, ws.max_column + 1):
        day = _parse_day_number(ws.cell(row=header_row, column=col).value)
        if day is not None:
            day_run_cols.append(col)

    if not day_run_cols:
        raise HeaderParseError(
            f"sheet {ws.title!r}：第 {header_row} 行找不到任何「N号」日期表头，"
            "识别不出平台区块——表头结构可能变了，需要人工确认"
        )

    blocks: list[tuple[int, int]] = []  # (起始列, 结束列)
    run_start = day_run_cols[0]
    prev_col = day_run_cols[0]
    for col in day_run_cols[1:]:
        if col != prev_col + 1:
            blocks.append((run_start, prev_col))
            run_start = col
        prev_col = col
    blocks.append((run_start, prev_col))

    platforms: dict[str, PlatformBlock] = {}
    for i, (min_col, max_col) in enumerate(blocks):
        day_columns: dict[int, int] = {}
        for col in range(min_col, max_col + 1):
            day = _parse_day_number(ws.cell(row=header_row, column=col).value)
            if day is None:
                raise HeaderParseError(
                    f"sheet {ws.title!r}：第 {header_row} 行第 {col} 列本来应该是"
                    f"「N号」这种日期表头（属于第 {min_col}-{max_col} 列的平台区块），"
                    f"实际内容是 {ws.cell(row=header_row, column=col).value!r}，识别不出是几号"
                )
            if day in day_columns:
                raise HeaderParseError(
                    f"sheet {ws.title!r}：第 {min_col}-{max_col} 列的平台区块里，「{day}号」"
                    f"这个日期出现了不止一列（第 {day_columns[day]} 列和第 {col} 列），分不清用哪一列"
                )
            day_columns[day] = col

        if i == 0:
            # 第一个区块固定是 WF-RX——这几个 sheet 里这个区块第1行的文字不统一
            # （空白/"Wayfair"/"Wayfair All SKU"），已经拿真实文件核对过，不是猜的。
            platform_key = "WF-RX"
        else:
            label = _label_at(ws, platform_label_row, min_col)
            norm = _norm_label(label)
            platform_key = _PLATFORM_ALIASES.get(norm)
            if platform_key is None:
                raise HeaderParseError(
                    f"sheet {ws.title!r}：第 {platform_label_row} 行、第 {min_col} 列附近的平台名"
                    f"{label!r} 不在已知的平台对照表里，需要人工确认这是哪个平台"
                )

        if platform_key in platforms:
            raise HeaderParseError(
                f"sheet {ws.title!r}：平台 {platform_key!r} 出现了不止一个区块，"
                "分不清该用哪一个"
            )
        platforms[platform_key] = PlatformBlock(platform_key=platform_key, day_columns=day_columns)

    sku_rows: dict[str, int] = {}
    duplicates: dict[str, list[int]] = {}
    for row in range(header_row + 1, ws.max_row + 1):
        value = ws.cell(row=row, column=sku_col).value
        if value is None or str(value).strip() == "":
            continue
        sku = str(value).strip()
        if sku in sku_rows:
            duplicates.setdefault(sku, [sku_rows[sku]]).append(row)
        else:
            sku_rows[sku] = row

    if duplicates:
        detail = "; ".join(f"{sku} 出现在第 {rows} 行" for sku, rows in duplicates.items())
        raise HeaderParseError(
            f"sheet {ws.title!r}：以下 SKU 在这个 sheet 里出现了不止一行，"
            f"分不清该写哪一行，需要人工先处理掉重复行：{detail}"
        )

    return SheetHeaderMap(sheet_name=ws.title, sku_column=sku_col, sku_rows=sku_rows, platforms=platforms)
