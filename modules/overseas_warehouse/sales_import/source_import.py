"""读 ERP 导出表 + CastleGate(CG) 平台 CSV 导出，转成统一的 (仓库, 平台, SKU, 数量) 记录，
再按 (仓库, 平台, SKU) 汇总求和。

ERP 导出文件后缀是 .xls，但实际内容是 xlsx（zip）格式——不是旧版二进制 xls，openpyxl
本来能读，只是它会先看文件后缀，看到 .xls 就直接报错拒绝，不会真的去看内容。绕过办法是
用 open(path, "rb") 传文件对象而不是路径字符串给 load_workbook——openpyxl 只有拿到路径
字符串时才会做后缀检查，拿到文件对象就会直接尝试当 zip 打开，见 openpyxl.reader.excel
的 _validate_archive。如果真的是旧版二进制 xls（万一以后 ERP 导出格式变了），zip 解析
会失败，这里会包装成更清楚的报错信息而不是让 openpyxl 的 BadZipFile 直接抛出去。

单行数据缺 SKU/数量 vs. 认不出「店铺」字段的值，处理方式不一样——不是随便挑的：
- 单独一行缺 SKU 或数量，是"这一笔订单本身有问题"（拿真实的 CastleGate CSV 导出验证过，
  确实会出现——一笔正常发货、价格/运单号都有的订单，Item Number/SKU 两个字段却是空的，
  大概率是平台那边同步出的问题），只影响这一行，跳过、记下来给人工事后核对就够了，不该
  因为一笔订单的问题让整份文件（可能几百行）全部导入失败。
- 「店铺」字段的值不认识（比如以后平台改了店铺命名），说明是我们的平台对照表本身跟不上了，
  很可能不止一行受影响——这种要直接报错，不能囫囵跳过，不然可能悄悄漏掉一大批数据自己都
  不知道。
"""
from __future__ import annotations

import csv
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from zipfile import BadZipFile

import openpyxl

from .platform_rules import (
    PLATFORM_OS,
    WAREHOUSE_CG,
    UnknownShopError,
    classify_erp_platform,
    classify_erp_warehouse,
)

ERP_REQUIRED_COLUMNS = ("采购SKU", "店铺", "发货仓库", "采购SKU个数")
CSV_REQUIRED_COLUMNS = ("Item Number", "Quantity")


class SourceReadError(Exception):
    pass


@dataclass(frozen=True)
class SourceRecord:
    warehouse: str
    platform: str
    sku: str
    quantity: int
    origin: str  # 出处，用于报错/未匹配清单里指给人工看，比如 "ERP 第5行"


@dataclass(frozen=True)
class ReadResult:
    records: list[SourceRecord]
    skipped: list[str] = field(default_factory=list)  # 人工可读的"跳过了哪一行、为什么"


def read_erp_export(path: str | Path) -> ReadResult:
    path = Path(path)
    # 必须全程都在 `with open(...)` 里面读完，不能提前把文件对象关掉——非 read_only 模式
    # 虽然是一次性整个解析进内存的，但保持跟旧代码一样的写法更保险，不用再纠结这件事。
    #
    # 这里故意不用 read_only=True：拿真实的 0915 那份 ERP 导出验证过，read_only 模式下
    # ws.max_row/max_column 是直接读 sheet XML 里 <dimension> 标签的声明值，不会真的去扫
    # <sheetData>——而这份文件的 <dimension> 标签写的是 "A1"（骗人的，可能是 ERP 导出
    # 工具自己的 bug），哪怕 <sheetData> 里其实有 71 行 48 列的真实数据，read_only 模式
    # 下 ws.max_row 还是会算成 1，导致后面读表头直接读到一个空表头、误判成"表头缺列"。
    # 非 read_only 模式会完整解析 <sheetData>、自己数出真实的行列数，不受这个假 dimension
    # 标签影响。ERP 导出文件体量很小（几十到几百行），完整解析也很快，不用担心性能。
    try:
        with open(path, "rb") as fh:
            wb = openpyxl.load_workbook(fh, read_only=False, data_only=True)
            return _read_erp_rows(wb, path)
    except BadZipFile as exc:
        raise SourceReadError(
            f"打不开 ERP 导出文件 {path.name}：不是能识别的 Excel 格式（{exc}）"
        ) from exc


def _read_erp_rows(wb, path: Path) -> ReadResult:
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        raise SourceReadError(f"ERP 导出文件 {path.name} 是空的，没有表头")

    col_index = {name: i for i, name in enumerate(header) if name is not None}
    missing = [c for c in ERP_REQUIRED_COLUMNS if c not in col_index]
    if missing:
        raise SourceReadError(
            f"ERP 导出文件 {path.name} 缺少必须的表头列：{missing}（现有表头：{list(header)}）"
        )

    records: list[SourceRecord] = []
    skipped: list[str] = []
    for row_num, row in enumerate(rows, start=2):
        sku = row[col_index["采购SKU"]]
        shop = row[col_index["店铺"]]
        ship_warehouse = row[col_index["发货仓库"]]
        qty = row[col_index["采购SKU个数"]]

        if sku is None and shop is None and ship_warehouse is None:
            continue  # 整行空白，跳过

        origin = f"ERP 第{row_num}行"
        if not sku or str(sku).strip() == "":
            skipped.append(f"{origin}：「采购SKU」是空的，跳过")
            continue
        if qty is None:
            skipped.append(f"{origin}：「采购SKU个数」是空的，跳过")
            continue

        try:
            platform = classify_erp_platform(shop)
        except UnknownShopError as exc:
            raise SourceReadError(f"{origin}：{exc}") from exc
        warehouse = classify_erp_warehouse(ship_warehouse)
        if warehouse == WAREHOUSE_CG:
            # CG 仓库自己的分平台数据是靠 CastleGate 那 3 份 CSV（RQ/RX/TS）导入的，ERP 里
            # 判定成 CG 的行不管「店铺」字段实际写的是什么，业务上统一都算 OS——这条只对
            # CG 生效，TX1/IL/CA1 的 ERP 行还是按「店铺」字段分到各自的平台。
            platform = PLATFORM_OS

        records.append(
            SourceRecord(
                warehouse=warehouse,
                platform=platform,
                sku=str(sku).strip(),
                quantity=int(qty),
                origin=origin,
            )
        )
    return ReadResult(records=records, skipped=skipped)


def read_castlegate_csv(path: str | Path, platform: str) -> ReadResult:
    """CG 的 CSV 导出，仓库固定是 CG，平台由调用方传进来（界面上人工选/确认过的）。"""
    path = Path(path)
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        if reader.fieldnames is None:
            raise SourceReadError(f"CSV 文件 {path.name} 是空的，没有表头")
        missing = [c for c in CSV_REQUIRED_COLUMNS if c not in reader.fieldnames]
        if missing:
            raise SourceReadError(
                f"CSV 文件 {path.name} 缺少必须的表头列：{missing}（现有表头：{reader.fieldnames}）"
            )

        records: list[SourceRecord] = []
        skipped: list[str] = []
        for row_num, row in enumerate(reader, start=2):
            sku = row.get("Item Number")
            qty = row.get("Quantity")
            if (sku is None or sku.strip() == "") and (qty is None or qty.strip() == ""):
                continue  # 空行跳过

            origin = f"{path.name} 第{row_num}行"
            if not sku or sku.strip() == "":
                skipped.append(f"{origin}：「Item Number」是空的，跳过")
                continue
            if not qty or qty.strip() == "":
                skipped.append(f"{origin}：「Quantity」是空的，跳过")
                continue

            try:
                qty_int = int(float(qty))
            except ValueError:
                skipped.append(f"{origin}：「Quantity」的值 {qty!r} 不是数字，跳过")
                continue

            records.append(
                SourceRecord(
                    warehouse="CG",
                    platform=platform,
                    sku=sku.strip(),
                    quantity=qty_int,
                    origin=origin,
                )
            )
    return ReadResult(records=records, skipped=skipped)


def aggregate_records(records: list[SourceRecord]) -> dict[tuple[str, str, str], int]:
    """按 (仓库, 平台, SKU) 汇总求和——同一天同一仓库同一平台同一 SKU 如果有多笔订单，
    数量直接加总写进那一天的格子。"""
    totals: dict[tuple[str, str, str], int] = defaultdict(int)
    for r in records:
        totals[(r.warehouse, r.platform, r.sku)] += r.quantity
    return dict(totals)
