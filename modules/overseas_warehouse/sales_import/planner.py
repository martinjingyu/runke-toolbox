"""把"源数据记录（仓库+平台+SKU+数量）"跟"目标汇总表当前的表头/数值"对一遍，
算出一份"哪个格子从什么值改成什么值"的改动预览，人工确认之后才真的写盘。

跟项目里其它涉及真实业务数据的工具（shipment_plan_apply 那几个）一个原则：先算出
before/after 给人看，不直接改。区别是这里最后落盘不能走 openpyxl（会把 WPS 单元格图片
弄丢，见 xlsx_writer.py 开头的说明），所以 apply_plan 用的是 core.backup 里新加的
atomic_replace_with_backup + 自己的 xlsx_writer，不是 core.backup.atomic_save_with_backup
（那个是给 openpyxl 的 wb.save() 用的）。
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from core.backup import atomic_replace_with_backup
from .header_map import HeaderParseError, SheetHeaderMap, parse_sheet_header
from .platform_rules import IMPORTABLE_PLATFORMS, WAREHOUSE_CA1, WAREHOUSE_CG, WAREHOUSE_IL, WAREHOUSE_TX1
from .source_import import SourceRecord
from .xlsx_writer import apply_cell_updates

# 汇总表实际的 sheet 名跟"仓库代号"不完全一样——CA1 那个 sheet 名字带个尾随空格
# （'CA1 '），这里不硬编码这个空格，运行时按"去空白后是不是一样"去匹配真实的 sheet 名，
# 防止哪天有人把这个空格清掉了代码就跟着失效。
_WAREHOUSE_KEYS = (WAREHOUSE_TX1, WAREHOUSE_IL, WAREHOUSE_CA1, WAREHOUSE_CG)


class PlanError(Exception):
    pass


@dataclass(frozen=True)
class DiffRow:
    warehouse: str
    sheet_name: str
    platform: str
    sku: str
    product_name: str
    row: int
    col: int
    old_value: float
    new_value: float


@dataclass(frozen=True)
class UnmatchedTotal:
    warehouse: str
    platform: str
    sku: str
    quantity: int
    reason: str  # 给人看的原因，比如"这个仓库的 sheet 不在了"/"这个 SKU 在表里找不到对应的行"


@dataclass(frozen=True)
class FuzzyMatchNote:
    """SKU 原文在汇总表里找不到，去掉最后一个"-"后面的部分再试了一次、匹配上了——
    不算"未匹配"，但也不能悄悄地就当没事发生，得让人知道"这批数据其实是按哪个 SKU 写进去的"。
    """
    warehouse: str
    platform: str
    source_sku: str
    matched_sku: str
    quantity: int
    origins: list[str]  # 具体来自哪个源文件的第几行，比如 "CG-RX export.csv 第141行"


@dataclass(frozen=True)
class ImportPlan:
    day: int
    target_path: str
    diff_rows: list[DiffRow]
    unmatched: list[UnmatchedTotal]
    fuzzy_matches: list[FuzzyMatchNote]


def _resolve_sheet_names(wb) -> dict[str, str]:
    """4 个仓库 sheet 不要求全都在——业务确认过，TX1 这种"只有库存"的仓库，卖完之后
    对应的 sheet 以后会被直接删掉，不是异常状态。缺了哪个仓库的 sheet，这个仓库的源数据
    到时候会走"未匹配"清单（build_plan 里判断），不会因为少一个 sheet 就让整批导入失败。

    但 4 个一个都不在，大概率是选错了文件（比如选到了别的表），这种情况还是要报错，
    不能悄悄地把所有源数据都塞进"未匹配"清单让人自己去猜为什么全都没匹配上。
    """
    by_stripped = {name.strip(): name for name in wb.sheetnames}
    resolved = {key: by_stripped[key] for key in _WAREHOUSE_KEYS if key in by_stripped}
    if not resolved:
        raise PlanError(
            f"目标表里一个仓库 sheet（{'/'.join(_WAREHOUSE_KEYS)}）都找不到，"
            f"是不是选错文件了？现有 sheet：{wb.sheetnames}"
        )
    return resolved


def _find_product_name_column(ws) -> int | None:
    for row in range(1, 6):
        for col in range(1, min(ws.max_column, 40) + 1):
            if ws.cell(row=row, column=col).value == "品名":
                return col
    return None


def _strip_last_dash_suffix(sku: str) -> str | None:
    """去掉 SKU 最后一个"-"后面的部分，比如"TD-332-RX" -> "TD-332"、"TD-794-1" -> "TD-794"
    ——各平台导出的 SKU 有时候会在真正的货号后面自己加一段后缀（平台代码、变体序号……），
    汇总表里登记的还是不带这段后缀的原始货号。

    只切一刀，不递归再切：像"TD-1"这种本来就是"货号-序号"基本形态的 SKU，切一刀会变成
    "TD"，明显是把货号本身切坏了，所以要求切完剩下的部分本身还得带"-"（还长得像
    "货号-序号"）才认为这一刀切得有意义，否则不做这个兜底，直接算未匹配。
    """
    idx = sku.rfind("-")
    if idx <= 0:
        return None
    prefix = sku[:idx]
    if "-" not in prefix:
        return None
    return prefix


def build_plan(
    target_path: str | Path,
    day: int,
    records: list[SourceRecord],
    clear_untouched_rows: bool = True,
) -> ImportPlan:
    """clear_untouched_rows=True（默认，业务明确要求的行为）：不管这次源数据实际覆盖了
    哪些"仓库+平台"，这一天在全部 4 个仓库 sheet × 全部 4 个平台（OS/WF-RX/WF-RQ/WF-TS）
    这一列，都会整列处理——这次有数据的 SKU 写新数据，没数据的 SKU（不管这个"仓库+平台"
    这次是不是提供了源文件）一律清成 0。

    这意味着**每次导入都必须把当天全部源文件（ERP + CG 的 3 份 CSV）一起选上**，漏选
    哪一份，那一份对应平台当天的真实销量会被一起清空，不会因为"这次没提供这个平台的
    文件"就保留它原来的值——这是业务确认过的、故意要的效果（宁可漏选时错误清空、也要
    保证不会有导入不掉的旧数据残留），不是遗漏考虑。
    """
    if not 1 <= day <= 31:
        raise PlanError(f"「几号」必须是 1-31 之间的数字，收到的是 {day!r}")

    target_path = Path(target_path)
    wb = openpyxl.load_workbook(target_path, read_only=False, data_only=True)
    sheet_names = _resolve_sheet_names(wb)

    header_maps: dict[str, SheetHeaderMap] = {}
    product_name_cols: dict[str, int | None] = {}

    def _get_header_map(warehouse_key: str) -> SheetHeaderMap:
        if warehouse_key not in header_maps:
            ws = wb[sheet_names[warehouse_key]]
            try:
                header_maps[warehouse_key] = parse_sheet_header(ws)
            except HeaderParseError as exc:
                raise PlanError(str(exc)) from exc
            product_name_cols[warehouse_key] = _find_product_name_column(ws)
        return header_maps[warehouse_key]

    # 按 (仓库, 平台, SKU原文) 先汇总一次，顺手把每一笔的出处（哪个源文件第几行）也留着——
    # 后面哪个 SKU 是靠"去掉后缀"这个兜底匹配上的，需要报给人看"是源文件里的哪个 SKU"。
    totals: dict[tuple[str, str, str], int] = defaultdict(int)
    origins_by_sku: dict[tuple[str, str, str], list[str]] = defaultdict(list)
    for r in records:
        key = (r.warehouse, r.platform, r.sku)
        totals[key] += r.quantity
        origins_by_sku[key].append(r.origin)

    unmatched: list[UnmatchedTotal] = []
    fuzzy_matches: list[FuzzyMatchNote] = []
    # 按"最终真的要写进哪一行"(仓库, 平台, 行号) 重新汇总——同一行可能被好几个不同写法的
    # 源 SKU（比如"TD-332"本身 + 兜底匹配上的"TD-332-RX"）一起摊上，数量要加在一起，
    # 不能後写的覆盖先写的。
    resolved: dict[tuple[str, str, int], dict] = {}

    for (warehouse, platform, sku), qty in sorted(totals.items()):
        if platform not in IMPORTABLE_PLATFORMS:
            # 理论上 source_import 只会产生这4个平台之一，这里是防御性检查。
            raise PlanError(f"内部错误：出现了不支持导入的平台 {platform!r}")
        if warehouse not in _WAREHOUSE_KEYS:
            raise PlanError(f"内部错误：出现了不认识的仓库代号 {warehouse!r}")

        if warehouse not in sheet_names:
            # 这个仓库的 sheet 目标表里没有（比如 TX1 卖完货之后 sheet 被删了）——不算错误，
            # 跟"SKU 在表里找不到"一样处理，进未匹配清单，不阻塞其它仓库正常导入。
            unmatched.append(UnmatchedTotal(
                warehouse=warehouse, platform=platform, sku=sku, quantity=qty,
                reason=f"目标表里没有「{warehouse}」这个仓库的 sheet（可能已经删掉了）",
            ))
            continue

        header_map = _get_header_map(warehouse)
        block = header_map.platforms.get(platform)
        if block is None:
            raise PlanError(
                f"「{sheet_names[warehouse]}」这个 sheet 里没有找到 {platform!r} 平台的表头区块，"
                "需要人工确认这张表是不是缺了这个平台的列"
            )
        col = block.day_columns.get(day)
        if col is None:
            raise PlanError(
                f"「{sheet_names[warehouse]}」的 {platform!r} 平台区块里没有「{day}号」这一列"
            )

        matched_sku = sku
        row = header_map.sku_rows.get(sku)
        if row is None:
            fallback_sku = _strip_last_dash_suffix(sku)
            if fallback_sku is not None:
                row = header_map.sku_rows.get(fallback_sku)
                if row is not None:
                    matched_sku = fallback_sku

        if row is None:
            unmatched.append(UnmatchedTotal(
                warehouse=warehouse, platform=platform, sku=sku, quantity=qty,
                reason=f"「{sheet_names[warehouse]}」里找不到这个 SKU 对应的行（去掉最后一段"
                       "后缀再试也没找到），需要人工核对是不是新品还没登记",
            ))
            continue

        if matched_sku != sku:
            fuzzy_matches.append(FuzzyMatchNote(
                warehouse=warehouse, platform=platform, source_sku=sku, matched_sku=matched_sku,
                quantity=qty, origins=origins_by_sku[(warehouse, platform, sku)],
            ))

        key = (warehouse, platform, row)
        if key not in resolved:
            resolved[key] = {"col": col, "matched_sku": matched_sku, "qty": 0}
        resolved[key]["qty"] += qty

    if clear_untouched_rows:
        # 不管这次源数据有没有覆盖到，目标表里现有的每个仓库 × 全部 4 个平台，这一天
        # 这一列都整列过一遍——业务明确要的效果，见函数开头的说明。
        for warehouse in sheet_names:
            header_map = _get_header_map(warehouse)
            ws = wb[sheet_names[warehouse]]
            for platform in IMPORTABLE_PLATFORMS:
                block = header_map.platforms.get(platform)
                if block is None or day not in block.day_columns:
                    continue  # 这个仓库的表里本来就没有这个平台的区块，没有列可清
                col = block.day_columns[day]
                for sku, row in header_map.sku_rows.items():
                    key = (warehouse, platform, row)
                    if key in resolved:
                        continue  # 这一行这次有真实数据，上面已经处理过
                    old_value = ws.cell(row=row, column=col).value or 0
                    if not old_value:
                        continue  # 本来就是空/0，没什么好清的
                    resolved[key] = {"col": col, "matched_sku": sku, "qty": 0}

    diff_rows: list[DiffRow] = []
    for (warehouse, platform, row), acc in sorted(resolved.items()):
        ws = wb[sheet_names[warehouse]]
        col = acc["col"]
        old_value = ws.cell(row=row, column=col).value or 0
        name_col = product_name_cols.get(warehouse)
        product_name = ws.cell(row=row, column=name_col).value if name_col else ""

        diff_rows.append(
            DiffRow(
                warehouse=warehouse,
                sheet_name=sheet_names[warehouse],
                platform=platform,
                sku=acc["matched_sku"],
                product_name=str(product_name or ""),
                row=row,
                col=col,
                old_value=old_value,
                new_value=acc["qty"],
            )
        )

    return ImportPlan(
        day=day, target_path=str(target_path), diff_rows=diff_rows,
        unmatched=unmatched, fuzzy_matches=fuzzy_matches,
    )


def apply_plan(plan: ImportPlan) -> Path:
    """把 plan 里的改动真的写进目标文件，写之前先备份。"""
    updates: dict[str, dict[tuple[int, int], float]] = {}
    for d in plan.diff_rows:
        updates.setdefault(d.sheet_name, {})[(d.row, d.col)] = d.new_value

    def _write(tmp_path: Path) -> None:
        apply_cell_updates(plan.target_path, str(tmp_path), updates)

    return atomic_replace_with_backup(plan.target_path, _write)
