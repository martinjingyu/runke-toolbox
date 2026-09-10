"""发货计划汇总表更新——这个工具的界面。

跟"发货计划自动更新"是同一套输入（同样读运营提供的发货计划表、同样先查在售产品信息总表
翻译成货号、同样按采购日期从早到晚分摊算出该扣哪几笔采购订单），但只写发货计划汇总表，
完全不碰采购订单汇总表——采购订单汇总表在这里只读、用来算"这个货号该分摊到哪几笔采购订单"
（发货计划汇总表待定行的匹配 key 是采购单号+型号，这两个字段来自分摊结果），不会被写入、
也不需要备份或存盘（见 planner.py 的 apply_plan_summary_only）。

不做预览、不弹二次确认——点一下「写入」就直接跑完整个流程，写入前仍然会先自动备份发货计划
汇总表原文件，这是唯一保留的安全网。如果这一批发货计划里有数据本身的问题（SKU 查不到货号、
解析报错……），在写入之前就会整批拦下来；"采购表里记录的量不够"是唯一的例外——只把那一行
跳过（不写、不算阻塞），其它行照常写入，跳过的行汇总成一张异常数据表格存在第一份发货计划表
旁边（见 planner.py 的 write_skipped_items_report）。这种"记录不够"可能是这一行当初就是
因为余量不够被「采购订单分摊更新」跳过了（正常），也可能是传错了采购表（数据对不上）——
分不清是哪种，见 build_plan_from_recorded_allocations 的说明；真传错了的话异常表格里会是
一整批"记录不够"，一眼就能看出来。
界面结构跟"发货计划自动更新"保持一致，只是少了采购订单汇总表那一栏的写入/备份/存盘流程。
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
from PySide6.QtCore import QDate, QSettings, Qt, QThread, Signal
from PySide6.QtWidgets import (
    QComboBox,
    QDateEdit,
    QFileDialog,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMessageBox,
    QProgressBar,
    QPushButton,
    QScrollArea,
    QTableWidget,
    QTableWidgetItem,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

from core.backup import backup_file

from .planner import apply_plan_summary_only, build_plan_from_recorded_allocations, write_skipped_items_report
from .product_lookup import load_product_lookup
from .purchase_book import PurchaseBook
from .shipment_summary import ShipmentSummaryBook
from .shipment_templates import PlanLine, list_sheet_names, parse_shipment_plan

_TEMPLATE_LABELS = {"walmart": "沃尔玛", "amazon": "亚马逊", "overseas": "海外仓"}
_LABEL_TO_TEMPLATE = {v: k for k, v in _TEMPLATE_LABELS.items()}
_AUTO_LABEL = "自动识别"

# 这两张表长期维护、路径基本不会变，记住上次选过的路径——跟另外两个工具分开存一份，
# 不共用同一个 key，避免这几个工具互相覆盖对方记的路径（往往选的是不同的文件）。
_SETTINGS_KEY_PRODUCT = "shipment_summary_apply/product_path"
_SETTINGS_KEY_PURCHASE = "shipment_summary_apply/purchase_path"
_SETTINGS_KEY_SUMMARY = "shipment_summary_apply/summary_path"


def _next_wednesday(today: dt.date | None = None) -> dt.date:
    today = today or dt.date.today()
    days_ahead = (2 - today.weekday()) % 7  # Monday=0 ... Wednesday=2
    return today + dt.timedelta(days=days_ahead)


def _file_picker_row(label_text: str, on_browse) -> tuple[QHBoxLayout, QLineEdit]:
    row = QHBoxLayout()
    row.addWidget(QLabel(label_text))
    line_edit = QLineEdit()
    line_edit.setReadOnly(True)
    row.addWidget(line_edit, 1)
    button = QPushButton("浏览…")
    button.clicked.connect(on_browse)
    row.addWidget(button)
    return row, line_edit


# 表格显示顺序也按这个排——处理时确实是按亚马逊>沃尔玛>海外仓分摊的（见 planner.py 的
# _TEMPLATE_PRIORITY），如果列表还是按运营原本添加文件的顺序摆，看着容易让人误以为是按
# 那个顺序处理，所以加文件之后顺手把表格也重排一遍，眼见为实。
_TEMPLATE_ORDER = {"amazon": 0, "walmart": 1, "overseas": 2}


def _entry_priority(entry: "_PlanFileEntry") -> int:
    return _TEMPLATE_ORDER.get(entry.effective_template_type, len(_TEMPLATE_ORDER))


class _PlanFileEntry:
    def __init__(self, path: str, sheet_names: list[str], detected: str | None):
        self.path = path
        self.sheet_names = sheet_names
        self.detected_template_type = detected  # 自动识别出来的模板——模板下拉框停在"自动
        # 识别"没被人工改过的时候，排序要按这个来，不能看下拉框当前文字（那时候文字就是
        # "自动识别"，映射不出具体模板）
        self.sheet_combo = QComboBox()
        self.sheet_combo.addItems(sheet_names)
        self.template_combo = QComboBox()
        self.template_combo.addItem(_AUTO_LABEL)
        self.template_combo.addItems(list(_TEMPLATE_LABELS.values()))
        if detected in _TEMPLATE_LABELS:
            self.template_combo.setCurrentText(_TEMPLATE_LABELS[detected])

    @property
    def sheet_name(self) -> str:
        return self.sheet_combo.currentText()

    @property
    def template_type(self) -> str | None:
        label = self.template_combo.currentText()
        return _LABEL_TO_TEMPLATE.get(label)

    @property
    def effective_template_type(self) -> str | None:
        # 人工在下拉框里明确选过的类型优先；没手动选过（还停在"自动识别"）就退回探测出来的
        # 那个，两边都拿不到才是真的不认识，排到最后。
        return self.template_type or self.detected_template_type

    def rebuild_combos(self, sheet_text: str, template_text: str) -> None:
        """重排表格时用——旧的 sheet_combo/template_combo 会随着 QTableWidget 清空整张表
        一起被销毁（setCellWidget 挂上去的控件归表格所有），这里重新建一份新的，把原来
        选的值（sheet_text/template_text）复原回去，不能指望旧控件还能继续用。
        """
        self.sheet_combo = QComboBox()
        self.sheet_combo.addItems(self.sheet_names)
        if sheet_text:
            self.sheet_combo.setCurrentText(sheet_text)
        self.template_combo = QComboBox()
        self.template_combo.addItem(_AUTO_LABEL)
        self.template_combo.addItems(list(_TEMPLATE_LABELS.values()))
        if template_text:
            self.template_combo.setCurrentText(template_text)


class _WriteWorker(QThread):
    blocked = Signal(list)
    backup_failed = Signal(str)
    save_failed = Signal(str, str)  # 报错信息, 发货计划汇总表备份文件名
    succeeded = Signal(str, int, str)  # 发货计划汇总表备份文件名, 跳过条数, 异常表格路径（没有就是空字符串）
    failed = Signal(str)
    stage = Signal(str)
    progress = Signal(str, int, int)

    def __init__(
        self,
        product_path: str,
        purchase_path: str,
        summary_path: str,
        plan_files: list[tuple[str, str, str | None]],
        ship_date: dt.date,
    ):
        super().__init__()
        self._product_path = product_path
        self._purchase_path = purchase_path
        self._summary_path = summary_path
        self._plan_files = plan_files
        self._ship_date = ship_date

    def run(self):
        try:
            self.stage.emit("正在读取在售产品信息总表…")
            lookup = load_product_lookup(Path(self._product_path))

            all_lines: list[PlanLine] = []
            parse_errors: list[str] = []
            for path, sheet_name, template_type in self._plan_files:
                self.stage.emit(f"正在解析「{Path(path).name}」…")
                parsed = parse_shipment_plan(Path(path), sheet_name, template_type)
                file_label = Path(path).name
                for line in parsed.lines:
                    line.source_file = file_label
                all_lines.extend(parsed.lines)
                parse_errors.extend(f"[{file_label}] {e}" for e in parsed.errors)

            # 采购订单汇总表这里只读、用来重放出"这条发货计划该冲抵哪笔采购订单"，不会被
            # 写入也不需要保存——按设计，传进来的这份表应该已经是「采购订单分摊更新」写过
            # 这一批之后的结果（见 build_plan_from_recorded_allocations 的说明），不是拿来
            # 重新按余量算一遍。
            self.stage.emit("正在加载采购订单汇总表…")
            purchase_wb = openpyxl.load_workbook(self._purchase_path, data_only=False)
            purchase_book = PurchaseBook(
                purchase_wb.active,
                progress_callback=lambda done, total: self.progress.emit(
                    "正在加载采购订单汇总表", done, total
                ),
            )

            self.stage.emit("正在加载发货计划汇总表…")
            summary_wb = openpyxl.load_workbook(self._summary_path, data_only=False)
            # 建待定库存索引这一步是自己按行扫的，能报真实的百分比，单独给一个阶段名，
            # 不然界面会一直停在"正在加载"不动，看起来像卡住了（见 panel.py 同样的说明）。
            self.stage.emit("正在扫描待定库存索引…")
            summary_book = ShipmentSummaryBook(
                summary_wb.active,
                progress_callback=lambda done, total: self.progress.emit(
                    "正在扫描待定库存索引", done, total
                ),
            )

            self.stage.emit("正在从采购订单汇总表重放分摊结果…")
            plan = build_plan_from_recorded_allocations(
                all_lines, parse_errors, lookup, purchase_book, self._ship_date
            )

            if plan.has_blocking_errors:
                lines = list(plan.parse_errors)
                for item in plan.items:
                    lines.extend(item.errors)
                self.blocked.emit(lines)
                return

            self.stage.emit("正在写入变化…")
            apply_plan_summary_only(
                plan,
                summary_book,
                progress_callback=lambda done, total: self.progress.emit("正在写入变化", done, total),
            )

            self.stage.emit("正在备份原文件…")
            try:
                summary_backup = backup_file(self._summary_path)
            except Exception as exc:
                self.backup_failed.emit(str(exc))
                return

            try:
                self.stage.emit("正在存盘发货计划汇总表…")
                summary_wb.save(self._summary_path)
            except Exception as exc:
                self.save_failed.emit(str(exc), summary_backup.name)
                return

            # 异常表格是"锦上添花"的附加产物，不是这次写入成不成功的一部分——文件都已经存盘了，
            # 这一步哪怕出问题也不该把整次操作报成失败，只是异常表格没生成而已。
            report_path = ""
            if plan.skipped_items:
                first_plan_path = Path(self._plan_files[0][0])
                try:
                    written = write_skipped_items_report(plan.skipped_items, first_plan_path)
                    report_path = str(written) if written is not None else ""
                except Exception:
                    report_path = ""

            self.succeeded.emit(summary_backup.name, len(plan.skipped_items), report_path)
        except Exception as exc:
            self.failed.emit(str(exc))


class ShipmentSummaryApplyPanel(QWidget):
    def __init__(self):
        super().__init__()
        self._plan_entries: list[_PlanFileEntry] = []
        self._worker: _WriteWorker | None = None
        self._settings = QSettings()

        outer_layout = QVBoxLayout(self)
        outer_layout.setContentsMargins(0, 0, 0, 0)
        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        scroll_area.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        outer_layout.addWidget(scroll_area)

        content = QWidget()
        scroll_area.setWidget(content)
        layout = QVBoxLayout(content)

        title = QLabel("发货计划汇总表更新")
        title.setStyleSheet("font-size: 16px; font-weight: bold;")
        layout.addWidget(title)
        note = QLabel(
            "只会修改发货计划汇总表，不碰采购订单汇总表——点「写入」直接生效，不会再弹预览确认。"
            "写入前会自动先备份发货计划汇总表原文件；这一批发货计划有数据问题的话，会在写入前"
            "整批拦下来。重放时按亚马逊>沃尔玛>海外仓的顺序处理，哪一行在采购表里记录的量不够"
            "（比如当初就是因为余量不够被跳过的）会被单独跳过（不算整批出错），跳过的行会汇总"
            "成一张异常数据表格。"
        )
        note.setWordWrap(True)
        layout.addWidget(note)

        inputs_box = QGroupBox("长期维护的两张表")
        inputs_layout = QVBoxLayout(inputs_box)
        row, self._product_edit = _file_picker_row("在售产品信息总表", self._browse_product)
        inputs_layout.addLayout(row)
        row, self._purchase_edit = _file_picker_row("采购订单汇总表（只读，用来算分摊）", self._browse_purchase)
        inputs_layout.addLayout(row)
        row, self._summary_edit = _file_picker_row("发货计划汇总表", self._browse_summary)
        inputs_layout.addLayout(row)
        layout.addWidget(inputs_box)

        self._product_edit.setText(self._settings.value(_SETTINGS_KEY_PRODUCT, ""))
        self._purchase_edit.setText(self._settings.value(_SETTINGS_KEY_PURCHASE, ""))
        self._summary_edit.setText(self._settings.value(_SETTINGS_KEY_SUMMARY, ""))

        plan_box = QGroupBox("运营提供的发货计划表（可以一次导入好几份）")
        plan_layout = QVBoxLayout(plan_box)
        add_row = QHBoxLayout()
        add_button = QPushButton("添加文件…")
        add_button.clicked.connect(self._add_plan_files)
        add_row.addWidget(add_button)
        add_row.addStretch(1)
        plan_layout.addLayout(add_row)

        self._plan_table = QTableWidget(0, 4)
        self._plan_table.setHorizontalHeaderLabels(["文件", "选哪个 sheet", "模板类型", ""])
        self._plan_table.horizontalHeader().setStretchLastSection(False)
        self._plan_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        plan_layout.addWidget(self._plan_table)
        layout.addWidget(plan_box)

        date_row = QHBoxLayout()
        date_row.addWidget(QLabel("这次发货日期（发货计划汇总表写这一天的记录）"))
        self._date_edit = QDateEdit()
        self._date_edit.setCalendarPopup(True)
        default_date = _next_wednesday()
        self._date_edit.setDate(QDate(default_date.year, default_date.month, default_date.day))
        date_row.addWidget(self._date_edit)
        date_row.addStretch(1)
        layout.addLayout(date_row)

        run_row = QHBoxLayout()
        self._write_button = QPushButton("写入")
        self._write_button.clicked.connect(self._start_write)
        run_row.addWidget(self._write_button)
        run_row.addStretch(1)
        layout.addLayout(run_row)

        self._progress = QProgressBar()
        self._progress.setRange(0, 0)
        self._progress.hide()
        layout.addWidget(self._progress)

        self._status_label = QLabel("")
        self._status_label.setWordWrap(True)
        layout.addWidget(self._status_label)

        self._error_text = QTextEdit()
        self._error_text.setReadOnly(True)
        self._error_text.hide()
        self._error_text.setMaximumHeight(160)
        layout.addWidget(self._error_text)

    # ---- 文件选择 ----

    def _browse_product(self):
        start_dir = str(Path(self._product_edit.text()).parent) if self._product_edit.text() else ""
        path, _ = QFileDialog.getOpenFileName(self, "选择在售产品信息总表", start_dir, "Excel 文件 (*.xlsx)")
        if path:
            self._product_edit.setText(path)
            self._settings.setValue(_SETTINGS_KEY_PRODUCT, path)

    def _browse_purchase(self):
        start_dir = str(Path(self._purchase_edit.text()).parent) if self._purchase_edit.text() else ""
        path, _ = QFileDialog.getOpenFileName(self, "选择采购订单汇总表", start_dir, "Excel 文件 (*.xlsx)")
        if path:
            self._purchase_edit.setText(path)
            self._settings.setValue(_SETTINGS_KEY_PURCHASE, path)

    def _browse_summary(self):
        start_dir = str(Path(self._summary_edit.text()).parent) if self._summary_edit.text() else ""
        path, _ = QFileDialog.getOpenFileName(self, "选择发货计划汇总表", start_dir, "Excel 文件 (*.xlsx)")
        if path:
            self._summary_edit.setText(path)
            self._settings.setValue(_SETTINGS_KEY_SUMMARY, path)

    def _add_plan_files(self):
        paths, _ = QFileDialog.getOpenFileNames(self, "选择运营提供的发货计划表", "", "Excel 文件 (*.xlsx)")
        for path in paths:
            try:
                sheet_names = list_sheet_names(Path(path))
                detected = None
                if sheet_names:
                    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                    from .shipment_templates import detect_template_type

                    detected = detect_template_type(wb[sheet_names[0]])
            except Exception as exc:
                QMessageBox.warning(self, "读取失败", f"「{Path(path).name}」读取失败：{exc}")
                continue

            entry = _PlanFileEntry(path, sheet_names, detected)
            self._plan_entries.append(entry)

        self._resort_plan_table()

    def _resort_plan_table(self) -> None:
        # 按亚马逊>沃尔玛>海外仓重排表格显示顺序，跟实际分摊处理的顺序看齐（见
        # _TEMPLATE_ORDER 的说明）。旧控件会随着 setRowCount(0) 被销毁，先按 entry 对象本身
        # 记下每个人当前选的 sheet/模板（必须在排序、清表之前记，晚了控件就没了），重排完
        # 再用 rebuild_combos() 建一份新的、把选择复原。
        selections_by_entry = {
            id(entry): (entry.sheet_name, entry.template_combo.currentText()) for entry in self._plan_entries
        }
        self._plan_entries.sort(key=_entry_priority)

        self._plan_table.setRowCount(0)
        for entry in self._plan_entries:
            sheet_text, template_text = selections_by_entry[id(entry)]
            entry.rebuild_combos(sheet_text, template_text)
            self._append_plan_row(entry)

    def _append_plan_row(self, entry: _PlanFileEntry) -> None:
        row = self._plan_table.rowCount()
        self._plan_table.insertRow(row)
        self._plan_table.setItem(row, 0, QTableWidgetItem(Path(entry.path).name))
        self._plan_table.setCellWidget(row, 1, entry.sheet_combo)
        self._plan_table.setCellWidget(row, 2, entry.template_combo)
        remove_button = QPushButton("移除")
        remove_button.clicked.connect(lambda: self._remove_plan_entry(entry))
        self._plan_table.setCellWidget(row, 3, remove_button)
        self._plan_table.resizeColumnsToContents()

    def _remove_plan_entry(self, entry: _PlanFileEntry) -> None:
        if entry not in self._plan_entries:
            return
        idx = self._plan_entries.index(entry)
        self._plan_entries.pop(idx)
        self._plan_table.removeRow(idx)

    # ---- 写入 ----

    def _start_write(self):
        product_path = self._product_edit.text().strip()
        purchase_path = self._purchase_edit.text().strip()
        summary_path = self._summary_edit.text().strip()

        missing = []
        if not product_path:
            missing.append("在售产品信息总表")
        if not purchase_path:
            missing.append("采购订单汇总表")
        if not summary_path:
            missing.append("发货计划汇总表")
        if not self._plan_entries:
            missing.append("发货计划表（至少一份）")
        if missing:
            QMessageBox.warning(self, "缺少输入", "还没选：" + "、".join(missing))
            return

        qdate = self._date_edit.date()
        ship_date = dt.date(qdate.year(), qdate.month(), qdate.day())

        plan_files = [(e.path, e.sheet_name, e.template_type) for e in self._plan_entries]

        self._write_button.setEnabled(False)
        self._progress.setRange(0, 0)
        self._progress.show()
        self._error_text.hide()
        self._status_label.setText("正在处理……表格行数多的话可能要几分钟，请耐心等待，不要关闭窗口。")

        self._worker = _WriteWorker(product_path, purchase_path, summary_path, plan_files, ship_date)
        self._worker.blocked.connect(self._on_blocked)
        self._worker.backup_failed.connect(self._on_backup_failed)
        self._worker.save_failed.connect(self._on_save_failed)
        self._worker.succeeded.connect(self._on_succeeded)
        self._worker.failed.connect(self._on_failed)
        self._worker.stage.connect(self._on_stage)
        self._worker.progress.connect(self._on_progress)
        self._worker.start()

    def _on_stage(self, text: str):
        self._progress.setRange(0, 0)
        self._status_label.setText(text)

    def _on_progress(self, stage: str, done: int, total: int):
        if total <= 0:
            return
        if self._progress.maximum() != total:
            self._progress.setRange(0, total)
        self._progress.setValue(done)
        self._status_label.setText(f"{stage}：{done}/{total}")

    def _on_blocked(self, error_lines: list):
        self._progress.hide()
        self._write_button.setEnabled(True)
        self._status_label.setText(f"这一批有 {len(error_lines)} 处问题，全部列在下面，一个都没写入。")
        self._error_text.setPlainText("\n".join(error_lines))
        self._error_text.show()

    def _on_succeeded(self, summary_backup_name: str, skipped_count: int, report_path: str):
        self._progress.hide()
        text = f"已写入。备份文件：{summary_backup_name}"
        if skipped_count > 0:
            text += f"\n有 {skipped_count} 行因为采购表里记录的量不够被跳过，没有写入。"
            text += f"\n异常数据表格：{report_path}" if report_path else "\n（异常数据表格生成失败，跳过的记录只能自己核对）"
        self._status_label.setText(text)
        self._write_button.setText("已写入")

    def _on_failed(self, message: str):
        self._progress.hide()
        self._write_button.setEnabled(True)
        self._status_label.setText("写入失败，见弹窗说明。")
        QMessageBox.critical(self, "写入失败", message)

    def _on_backup_failed(self, message: str):
        self._progress.hide()
        self._write_button.setEnabled(True)
        self._status_label.setText("写入失败，见弹窗说明。")
        QMessageBox.critical(self, "备份失败", f"没能先备份原文件，写入已取消：{message}")

    def _on_save_failed(self, message: str, summary_backup_name: str):
        self._progress.hide()
        self._write_button.setEnabled(True)
        self._status_label.setText("写入失败，见弹窗说明。")
        QMessageBox.critical(
            self,
            "写入失败",
            f"存盘失败：{message}\n\n"
            f"备份文件还在（{summary_backup_name}），原文件可能已经部分改动，建议手动检查。",
        )

    def stop_running_tasks(self):
        if self._worker is not None and self._worker.isRunning():
            self._worker.wait(30000)
