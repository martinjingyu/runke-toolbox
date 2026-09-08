"""读取发货计划表，只找「标签」在 wanted_labels 里的、「仓库」包含 warehouse_code、且
「状态=未发货」的行——不要求「发货时间=待定」，只要还没发货就算数，不管是已经安排了具体
发货日期还是仍然待定。

warehouse_code 由调用方传入（见 splitter.py 里从标签 PDF 文件名推导站点代号的说明），不在
这里写死成某一个站点——真实表里「仓库」的值形如"US(CA1)"，不是精确等于站点代号，所以按包含
匹配，不做精确相等。同一个标签完全可能同时发往好几个不同仓库（发货计划表里横跨全公司所有
站点），不加这个条件的话，箱数会把发去别的仓库、跟这份标签 PDF 毫无关系的待发记录也一起加
进来，数字对不上。

wanted_labels 由调用方传入，是标签 PDF 里实际出现过的那些标签——发货计划表本身有几万行，
横跨所有产品全部历史批次，跟这一份标签 PDF 完全无关的标签不该在这里被读出来，也不该被拿去
跟 PDF 做差集提示（不然会报出一大堆"计划里有、这份 PDF 没有"的噪音，那些标签根本不属于
这份 PDF，提示了也没意义）——先看这份标签 PDF 里到底有哪些标签，再拿着这份名单去计划表里逐个
核对，这才是这个工具该有的方向。

按「标签」分组、把这些行的「工厂」和「箱数」汇总起来，给拆分标签 PDF 用：
    标签 -> 这一批待发记录的工厂 + 箱数总和

同一个标签匹配到的待发行「工厂」不一致的话不硬选一个——这种情况单独记下来（见
PendingGroups.conflicts），交给调用方决定怎么提示人工核对，不在这里报错中断整个读取。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from ..shipment_plan_apply.column_utils import column_index_map, find_header_row, require_columns

REQUIRED_HEADERS = ["标签", "状态", "箱数", "工厂", "仓库"]

PENDING_STATUS = "未发货"


@dataclass
class PendingGroup:
    label: str
    factory: str
    total_boxes: float
    boxes_exact: bool
    row_count: int


@dataclass
class PendingGroups:
    groups: dict[str, PendingGroup] = field(default_factory=dict)
    conflicts: dict[str, list[str]] = field(default_factory=dict)  # 标签 -> 不一致的工厂列表


def _num(value) -> float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    raise ValueError(f"「箱数」列的值「{value}」不是数字")


def load_pending_groups(
    xlsx_path: str | Path, wanted_labels: set[str], warehouse_code: str, sheet_name: str | None = None
) -> PendingGroups:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]

    header_row = find_header_row(ws, REQUIRED_HEADERS, max_scan_rows=10)
    cols = column_index_map(ws, header_row)
    idx = require_columns(cols, REQUIRED_HEADERS, "发货计划表")

    factories: dict[str, set[str]] = {}
    boxes: dict[str, float] = {}
    counts: dict[str, int] = {}

    for row in ws.iter_rows(min_row=header_row + 1):
        label = row[idx["标签"] - 1].value
        if label is None:
            continue
        label = str(label).strip()
        if label not in wanted_labels:
            continue

        status = row[idx["状态"] - 1].value
        if status != PENDING_STATUS:
            continue

        warehouse = row[idx["仓库"] - 1].value
        if warehouse is None or warehouse_code not in str(warehouse):
            continue

        factory = row[idx["工厂"] - 1].value
        factory = str(factory).strip() if factory is not None else ""
        boxes_value = row[idx["箱数"] - 1].value

        factories.setdefault(label, set()).add(factory)
        boxes[label] = boxes.get(label, 0.0) + _num(boxes_value)
        counts[label] = counts.get(label, 0) + 1

    result = PendingGroups()
    for label, fs in factories.items():
        if len(fs) > 1:
            result.conflicts[label] = sorted(fs)
            continue
        total = boxes[label]
        exact = float(total).is_integer()
        result.groups[label] = PendingGroup(
            label=label,
            factory=next(iter(fs)),
            total_boxes=int(total) if exact else total,
            boxes_exact=exact,
            row_count=counts[label],
        )
    return result
