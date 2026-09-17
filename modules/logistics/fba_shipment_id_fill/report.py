"""把 FBA 数据里被跳过的行（日期不匹配/已取消/数量异常）汇总成一份说明表格，存在 FBA
数据文件夹旁边——跟 shipment_plan_apply/planner.py 的 write_skipped_items_report 是同一个
思路：不是错误，不阻塞写入，只是让人知道这一批里哪些行没有被处理、为什么。
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl

from .fba_source import SkippedFbaRow


def write_skip_report(skipped: list[SkippedFbaRow], folder: Path, duplicates_removed: int = 0) -> Path | None:
    if not skipped and not duplicates_removed:
        return None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "跳过说明"
    if duplicates_removed:
        ws.append([f"另有 {duplicates_removed} 条记录因为在多份 CSV 里完全重复出现，已自动去重（不算问题）"])
        ws.append([])
    ws.append(["来源文件", "表格行数", "货件名称", "MSKU", "跳过原因"])
    for row in skipped:
        ws.append([row.source_file, row.source_row, row.shipment_name, row.msku, row.reason])
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 22

    timestamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    base_name = f"FBA数据-跳过说明-{timestamp}"
    out_path = Path(folder) / f"{base_name}.xlsx"
    suffix = 2
    while out_path.exists():
        out_path = Path(folder) / f"{base_name}-{suffix}.xlsx"
        suffix += 1
    wb.save(out_path)
    return out_path
